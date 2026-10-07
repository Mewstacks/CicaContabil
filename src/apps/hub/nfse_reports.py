from __future__ import annotations

import io
import re
from collections.abc import Iterable
from datetime import date
from decimal import Decimal, InvalidOperation
from html import escape
from typing import Any, TypedDict

from django.core.exceptions import ObjectDoesNotExist
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import (  # type: ignore[import-untyped]
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.table import Table as WorksheetTable  # type: ignore[import-untyped]
from openpyxl.worksheet.table import TableStyleInfo
from reportlab.lib import colors  # type: ignore[import-untyped]
from reportlab.lib.enums import TA_CENTER, TA_RIGHT  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import A4, landscape  # type: ignore[import-untyped]
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # type: ignore[import-untyped]
from reportlab.lib.units import mm  # type: ignore[import-untyped]
from reportlab.platypus import (  # type: ignore[import-untyped]
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from apps.hub.models import NfseDocument

# Since NT SE/CGNFS-e 007/2026 the retained PIS, COFINS and CSLL come summed in vRetCSLL,
# so they are reported together as CRF (see apps.hub.nfse_facts).
TAXES = (
    ("iss", "ISS"),
    ("crf", "CRF (PIS/COFINS/CSLL)"),
    ("irrf", "IRRF"),
    ("inss", "INSS"),
)
SITUATION_LABELS = {"active": "Ativa", "cancelled": "Cancelada", "substituted": "Substituída"}


class RetentionReportRow(TypedDict):
    company: str
    dominio_code: str
    number: str
    direction: str
    direction_label: str
    issued_at: date | None
    competence: str
    counterparty: str
    service_code: str
    service_description: str
    amount: Decimal
    net_amount: Decimal
    situation: str
    counterparty_document: str
    retentions: dict[str, Decimal]
    retained_total: Decimal


def _decimal(value: object) -> Decimal:
    try:
        parsed = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0.00")
    if not parsed.is_finite() or parsed < 0:
        return Decimal("0.00")
    return parsed.quantize(Decimal("0.01"))


def _side_of(document: NfseDocument) -> Any:
    try:
        return document.side
    except ObjectDoesNotExist:
        return None


def _date_from_document(document: NfseDocument, normalized: dict[str, Any]) -> date | None:
    if document.issued_at:
        issued = document.issued_at
        return timezone.localtime(issued).date() if timezone.is_aware(issued) else issued.date()
    raw = normalized.get("issued_at")
    if not isinstance(raw, str):
        return None
    parsed_datetime = parse_datetime(raw)
    if parsed_datetime:
        if timezone.is_aware(parsed_datetime):
            parsed_datetime = timezone.localtime(parsed_datetime)
        return parsed_datetime.date()
    return parse_date(raw[:10])


def _safe_text(value: object, *, limit: int) -> str:
    return " ".join(str(value or "").split())[:limit]


def build_retention_report_rows(
    documents: Iterable[NfseDocument], *, demo: bool = False
) -> list[RetentionReportRow]:
    rows: list[RetentionReportRow] = []
    for document in documents:
        normalized = document.normalized_data if isinstance(document.normalized_data, dict) else {}
        side = _side_of(document)
        direction = str((side.direction if side else "") or normalized.get("direction") or "")
        if direction not in {"provided", "taken"}:
            digits = re.sub(r"\D", "", document.source_nsu) if demo else ""
            direction = (
                "provided"
                if digits and int(digits[-1]) % 2 == 0
                else "taken"
                if digits
                else "unknown"
            )
        direction_label = {
            "provided": "Saída · serviço prestado",
            "taken": "Entrada · serviço tomado",
            "unknown": "A confirmar",
        }[direction]
        has_facts = bool(side is not None and side.facts_version)
        if has_facts:
            retentions = {
                "iss": _decimal(side.iss_retained),
                "crf": _decimal(side.crf_retained),
                "irrf": _decimal(side.irrf_retained),
                "inss": _decimal(side.inss_retained),
            }
        else:
            raw = normalized.get("retentions")
            raw = raw if isinstance(raw, dict) else {}
            retentions = {
                "iss": _decimal(raw.get("iss")),
                "crf": sum(
                    (_decimal(raw.get(key)) for key in ("pis", "cofins", "csll")),
                    Decimal("0.00"),
                ),
                "irrf": _decimal(raw.get("irrf")),
                "inss": _decimal(raw.get("inss")),
            }
        retained_total = sum(retentions.values(), Decimal("0.00"))
        issued_at = (side.issued_on if has_facts else None) or _date_from_document(
            document, normalized
        )
        if has_facts and side.competence:
            competence = side.competence.strftime("%m/%Y")
        else:
            competence = _safe_text(normalized.get("competence"), limit=40)
            if not competence and issued_at:
                competence = issued_at.strftime("%m/%Y")
            elif len(competence) >= 7 and competence[4:5] == "-":
                competence = f"{competence[5:7]}/{competence[:4]}"
        amount = (
            _decimal(side.service_amount)
            if has_facts and side.service_amount is not None
            else _decimal(normalized.get("amount"))
        )
        rows.append(
            {
                "company": _safe_text(document.company.name, limit=160),
                "dominio_code": _safe_text(document.company.dominio_code, limit=40),
                "number": _safe_text(
                    (side.number if has_facts else "") or normalized.get("number"), limit=80
                )
                or "Sem número",
                "direction": direction,
                "direction_label": direction_label,
                "issued_at": issued_at,
                "competence": competence or "Não informada",
                "counterparty": _safe_text(
                    (side.counterparty_name if side else "") or normalized.get("counterparty_name"),
                    limit=160,
                )
                or ("Contraparte fictícia" if demo else "Não identificada"),
                "service_code": _safe_text(normalized.get("service_code"), limit=80),
                "service_description": _safe_text(normalized.get("service_description"), limit=500),
                "amount": amount,
                "net_amount": _decimal(side.net_amount if has_facts else None),
                "situation": SITUATION_LABELS.get(
                    str(side.situation if has_facts else "active"), "Ativa"
                ),
                "counterparty_document": side.counterparty_document if has_facts else "",
                "retentions": retentions,
                "retained_total": retained_total,
            }
        )
    return rows


def retention_totals(rows: Iterable[RetentionReportRow]) -> dict[str, Decimal]:
    totals = {key: Decimal("0.00") for key, _label in TAXES}
    totals["amount"] = Decimal("0.00")
    totals["retained_total"] = Decimal("0.00")
    for row in rows:
        totals["amount"] += row["amount"]
        totals["retained_total"] += row["retained_total"]
        for key, _label in TAXES:
            totals[key] += row["retentions"][key]
    return totals


def _brl(value: Decimal) -> str:
    number = f"{value:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return f"R$ {number}"


def generate_retention_pdf(rows: list[RetentionReportRow], *, demo: bool = False) -> bytes:
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("CicaBody", parent=styles["BodyText"], fontSize=8, leading=10)
    small = ParagraphStyle("CicaSmall", parent=body, fontSize=7, leading=8.5)
    right = ParagraphStyle("CicaRight", parent=small, alignment=TA_RIGHT)
    centered = ParagraphStyle("CicaCentered", parent=body, alignment=TA_CENTER)
    title = ParagraphStyle(
        "CicaTitle",
        parent=styles["Title"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#16352F"),
    )
    story: list[Any] = []
    generated_at = timezone.localtime(timezone.now())
    heading = (
        "Relatório demonstrativo de retenções NFS-e" if demo else "Relatório de retenções NFS-e"
    )
    story.append(Paragraph(heading, title))
    story.append(
        Paragraph(
            f"{len(rows)} nota{'s' if len(rows) != 1 else ''} · "
            f"gerado em {generated_at:%d/%m/%Y %H:%M}",
            centered,
        )
    )
    story.append(Spacer(1, 5 * mm))
    totals = retention_totals(rows)
    summary_data = [["Indicador", "Valor"], ["Valor das notas", _brl(totals["amount"])]]
    summary_data.extend([[f"{label} retido", _brl(totals[key])] for key, label in TAXES])
    summary_data.append(["Total retido", _brl(totals["retained_total"])])
    summary = Table(summary_data, colWidths=[60 * mm, 38 * mm], repeatRows=1)
    summary.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16352F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#DDEFE8")),
                ("ALIGN", (1, 1), (1, -1), "RIGHT"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7C8C2")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F4F7F6")]),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(KeepTogether([Paragraph("Resumo", styles["Heading2"]), summary]))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph("Notas do recorte", styles["Heading2"]))
    story.append(
        Paragraph(
            "Valores reproduzem somente retenções explícitas recebidas no XML. "
            "Ausência de valor não representa cálculo ou parecer fiscal.",
            small,
        )
    )
    story.append(Spacer(1, 2 * mm))
    detail: list[list[Any]] = [
        [
            "Empresa / nota",
            "Movimento",
            "Emissão / competência",
            "Contraparte / serviço",
            "Valor",
            "Retenções",
        ]
    ]
    for row in rows:
        taxes = [
            f"{label} {_brl(row['retentions'][key])}"
            for key, label in TAXES
            if row["retentions"][key] > 0
        ]
        taxes.append(f"Total {_brl(row['retained_total'])}")
        service = " · ".join(
            part for part in (row["service_code"], row["service_description"]) if part
        )
        detail.append(
            [
                Paragraph(f"<b>{escape(row['company'])}</b><br/>{escape(row['number'])}", small),
                Paragraph(escape(row["direction_label"]), small),
                Paragraph(
                    (row["issued_at"].strftime("%d/%m/%Y") if row["issued_at"] else "Não informada")
                    + f"<br/>{escape(row['competence'])}",
                    small,
                ),
                Paragraph(
                    f"<b>{escape(row['counterparty'])}</b>"
                    + (f"<br/>{escape(service[:160])}" if service else ""),
                    small,
                ),
                Paragraph(_brl(row["amount"]), right),
                Paragraph("<br/>".join(escape(item) for item in taxes), right),
            ]
        )
    table = Table(
        detail,
        colWidths=[44 * mm, 30 * mm, 31 * mm, 61 * mm, 27 * mm, 50 * mm],
        repeatRows=1,
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16352F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 7.5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C7D3CF")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7F6")]),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#5E6E69"))
        canvas.drawString(14 * mm, 8 * mm, "CICA · conferência operacional de NFS-e")
        canvas.drawRightString(283 * mm, 8 * mm, f"Página {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=heading,
        author="CICA",
    )
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def _xlsx_text(value: object) -> str:
    text = _safe_text(value, limit=1000)
    return f"'{text}" if text.lstrip().startswith(("=", "+", "-", "@")) else text


def generate_retention_xlsx(rows: list[RetentionReportRow], *, demo: bool = False) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    assert summary is not None
    summary.title = "Resumo"
    detail = workbook.create_sheet("Notas")
    dark = "16352F"
    pale = "DDEFE8"
    header_fill = PatternFill("solid", fgColor=dark)
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="C7D3CF")
    border = Border(bottom=thin)
    generated_at = timezone.localtime(timezone.now()).replace(tzinfo=None)
    last_detail_row = len(rows) + 1
    headers = [
        "Empresa",
        "Código Domínio",
        "Número da nota",
        "Situação",
        "Movimento",
        "Data de emissão",
        "Competência",
        "CNPJ/CPF contraparte",
        "Contraparte",
        "Código do serviço",
        "Descrição do serviço",
        "Valor do serviço",
        "Valor líquido",
        *[f"{label} retido" for _key, label in TAXES],
        "Total retido",
    ]
    service_column = headers.index("Valor do serviço") + 1
    first_tax = service_column + 2
    last_tax = first_tax + len(TAXES) - 1
    total_column = last_tax + 1
    service_letter = get_column_letter(service_column)
    last_letter = get_column_letter(total_column)
    summary_rows: list[list[object]] = [
        ["Relatório de retenções NFS-e" + (" · DEMONSTRAÇÃO" if demo else ""), None],
        ["Gerado em", generated_at],
        ["Notas no recorte", len(rows)],
        ["Valor das notas", f"=SUM(Notas!{service_letter}2:{service_letter}{last_detail_row})"],
    ]
    for index, (_key, label) in enumerate(TAXES):
        column = get_column_letter(first_tax + index)
        summary_rows.append([f"{label} retido", f"=SUM(Notas!{column}2:{column}{last_detail_row})"])
    summary_rows.append(
        ["Total retido", f"=SUM(Notas!{last_letter}2:{last_letter}{last_detail_row})"]
    )
    for values in summary_rows:
        summary.append(values)
    summary.merge_cells("A1:B1")
    summary["A1"].fill = header_fill
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 30
    summary["B2"].number_format = "dd/mm/yyyy hh:mm"
    for row_number in range(4, summary.max_row + 1):
        summary.cell(row_number, 2).number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
    for row_number in range(2, summary.max_row + 1):
        summary.cell(row_number, 1).font = Font(bold=True)
        summary.cell(row_number, 1).border = border
        summary.cell(row_number, 2).border = border
    summary.cell(summary.max_row, 1).fill = PatternFill("solid", fgColor=pale)
    summary.cell(summary.max_row, 2).fill = PatternFill("solid", fgColor=pale)
    summary.cell(summary.max_row, 1).font = Font(bold=True)
    summary.cell(summary.max_row, 2).font = Font(bold=True)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 22
    summary.freeze_panes = "A2"

    detail.append(headers)
    for cell in detail[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    first_tax_letter = get_column_letter(first_tax)
    last_tax_letter = get_column_letter(last_tax)
    for row in rows:
        detail.append(
            [
                _xlsx_text(row["company"]),
                _xlsx_text(row["dominio_code"]),
                _xlsx_text(row["number"]),
                row["situation"],
                row["direction_label"],
                row["issued_at"],
                _xlsx_text(row["competence"]),
                _xlsx_text(row["counterparty_document"]),
                _xlsx_text(row["counterparty"]),
                _xlsx_text(row["service_code"]),
                _xlsx_text(row["service_description"]),
                float(row["amount"]),
                float(row["net_amount"]),
                *[float(row["retentions"][key]) for key, _label in TAXES],
                None,
            ]
        )
        current = detail.max_row
        detail.cell(
            current,
            total_column,
            f"=SUM({first_tax_letter}{current}:{last_tax_letter}{current})",
        )
    for row_number in range(2, detail.max_row + 1):
        detail.cell(row_number, 6).number_format = "dd/mm/yyyy"
        for column in range(service_column, total_column + 1):
            detail.cell(row_number, column).number_format = "R$ #,##0.00;[Red]-R$ #,##0.00"
        for cell in detail[row_number]:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 11)
    widths = [30, 14, 16, 13, 26, 15, 13, 21, 32, 16, 44, 16, 16, 14, 22, 14, 14, 16]
    for index, width in enumerate(widths, 1):
        detail.column_dimensions[get_column_letter(index)].width = width
    detail.freeze_panes = "A2"
    if detail.max_row >= 2:
        table = WorksheetTable(
            displayName="RetencoesNFSe", ref=f"A1:{last_letter}{detail.max_row}"
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
        )
        detail.add_table(table)
    else:
        detail.auto_filter.ref = f"A1:{last_letter}1"
    detail.sheet_view.showGridLines = False
    summary.sheet_view.showGridLines = False
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
