"""Durable, explainable reconciliation intake and accounting services.

This module deliberately keeps parsing, classification and posting separate.  It is
safe for a job to execute more than once: source keys and database constraints are
the final idempotency boundary, not a Celery delivery guarantee.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from itertools import islice
from pathlib import Path, PurePath
from typing import Any, cast

from django.conf import settings
from django.core.files.base import ContentFile
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.utils import timezone

from apps.accounts.models import User
from apps.audit.services import record_event
from apps.hub.models import (
    AccountingExport,
    AccountingPeriod,
    JournalEntry,
    JournalLine,
    LedgerAccount,
    MovementReconciliation,
    NormalizedMovement,
    ReconciliationLayout,
    ReconciliationRule,
    ReconciliationRun,
    ReconciliationSourceFile,
)
from apps.organizations.models import Organization

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
MAX_PAGES = 500
OCR_REVIEW_THRESHOLD = 90
SAFE_SUFFIXES = {".ofx": "ofx", ".qfx": "ofx", ".csv": "csv", ".xlsx": "xlsx", ".pdf": "pdf"}
DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%y")


class ReconciliationError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedRecord:
    source_key: str
    occurred_on: date | None
    original_date: str
    amount_cents: int
    original_amount: str
    description: str
    document_number: str = ""
    counterparty: str = ""
    source_reference: dict[str, Any] | None = None
    confidence: int = 100
    confidence_details: dict[str, Any] | None = None


def detect_kind(filename: str, content: bytes) -> str:
    suffix = PurePath(filename).suffix.casefold()
    kind = SAFE_SUFFIXES.get(suffix)
    if kind is None:
        raise ReconciliationError("Use OFX, CSV, XLSX ou PDF.")
    if not content or len(content) > MAX_BYTES:
        raise ReconciliationError("Arquivo vazio ou maior que 25 MiB.")
    if kind == "pdf" and not content.startswith(b"%PDF-"):
        raise ReconciliationError("O conteúdo não corresponde a um PDF.")
    if kind == "xlsx" and not content.startswith(b"PK"):
        raise ReconciliationError("O conteúdo não corresponde a uma planilha XLSX.")
    if kind == "xlsx":
        _load_xlsx_workbook(content)
    if kind == "ofx" and b"OFX" not in content[:4096].upper():
        raise ReconciliationError("O conteúdo não corresponde a um OFX.")
    if kind == "csv" and b"\x00" in content:
        raise ReconciliationError("CSV contém dados binários inesperados.")
    return kind


def _tesseract_executable() -> str | None:
    """Resolve an explicitly configured binary or the standard local installer path."""
    configured = str(getattr(settings, "RECONCILIATION_TESSERACT_PATH", "") or "").strip()
    candidates = (
        configured,
        shutil.which("tesseract") or "",
        str(Path(os.environ.get("PROGRAMFILES", "")) / "Tesseract-OCR" / "tesseract.exe"),
    )
    return next(
        (candidate for candidate in candidates if candidate and Path(candidate).is_file()), None
    )


def _tesseract_environment() -> dict[str, str]:
    """Pass a private language-data directory to Tesseract when configured."""
    configured = str(getattr(settings, "RECONCILIATION_TESSDATA_PATH", "") or "").strip()
    local_default = Path(os.environ.get("LOCALAPPDATA", "")) / "CICA" / "tessdata"
    language_data = Path(configured) if configured else local_default
    environment = os.environ.copy()
    if language_data.is_dir():
        environment["TESSDATA_PREFIX"] = str(language_data)
    return environment


def local_ocr_available() -> bool:
    """Require both the local executable and Portuguese model before offering OCR."""
    tesseract_path = _tesseract_executable()
    if tesseract_path is None:
        return False
    try:
        result = subprocess.run(  # noqa: S603 - resolved local binary, fixed arguments
            [tesseract_path, "--list-langs"],
            capture_output=True,
            check=False,
            text=True,
            timeout=5,
            env=_tesseract_environment(),
        )
    except (OSError, subprocess.TimeoutExpired):
        return False
    return result.returncode == 0 and "por" in {line.strip() for line in result.stdout.splitlines()}


@transaction.atomic
def create_source_file(
    *,
    organization: Organization,
    company: Any,
    filename: str,
    content: bytes,
    origin: str,
    actor: object = None,
    physical_batch: str = "",
    financial_account: Any = None,
    request: object = None,
) -> tuple[ReconciliationSourceFile, ReconciliationRun, bool]:
    if company.organization_id != organization.id:
        raise ReconciliationError("A empresa não pertence ao escritório atual.")
    if financial_account is not None and (
        financial_account.organization_id != organization.id
        or financial_account.company_id != company.id
        or not financial_account.active
    ):
        raise ReconciliationError("A conta financeira não pertence à empresa ou está inativa.")
    kind = detect_kind(filename, content)
    digest = hashlib.sha256(content).hexdigest()
    existing = ReconciliationSourceFile.objects.filter(
        organization=organization, company=company, content_hash=digest
    ).first()
    if existing is not None:
        if financial_account is not None and existing.financial_account_id != financial_account.id:
            raise ReconciliationError(
                "Este arquivo já foi importado para outra conta financeira da empresa."
            )
        run = existing.runs.order_by("-created_at").first()
        if run is None:
            run = ReconciliationRun.objects.create(
                organization=organization,
                source_file=existing,
                created_by=actor if isinstance(actor, User) else None,
            )
        return existing, run, False
    source = ReconciliationSourceFile(
        organization=organization,
        company=company,
        financial_account=financial_account,
        kind=kind,
        origin=origin,
        original_filename=PurePath(filename).name[:255],
        content_hash=digest,
        content_type={
            "ofx": "application/x-ofx",
            "csv": "text/csv",
            "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "pdf": "application/pdf",
        }[kind],
        size_bytes=len(content),
        uploaded_by=actor if isinstance(actor, User) else None,
    )
    source.content.save(source.original_filename, ContentFile(content), save=False)
    try:
        # Keep the uniqueness boundary in PostgreSQL.  The inner transaction is
        # required so an IntegrityError from a simultaneous upload does not mark
        # the outer upload transaction unusable before we can reuse its source.
        with transaction.atomic():
            source.save()
    except IntegrityError:
        # The bytes are staged before the database row so private storage has a
        # deterministic name.  A concurrent request may win the unique hash
        # constraint; remove this losing staged object before reusing its row.
        source.content.delete(save=False)
        existing = ReconciliationSourceFile.objects.filter(
            organization=organization, company=company, content_hash=digest
        ).first()
        if existing is None:
            raise
        run = existing.runs.order_by("-created_at").first()
        if run is None:
            run = ReconciliationRun.objects.create(
                organization=organization,
                source_file=existing,
                created_by=actor if isinstance(actor, User) else None,
            )
        return existing, run, False
    layout = _compatible_layout_for_source(source)
    checkpoint: dict[str, Any] = {"physical_batch": physical_batch[:120]}
    if layout is not None:
        checkpoint["layout_selected_automatically"] = True
        checkpoint["layout_name"] = layout.name
        checkpoint["layout_version"] = layout.version
    run = ReconciliationRun.objects.create(
        organization=organization,
        source_file=source,
        layout_version=layout,
        checkpoint=checkpoint,
        created_by=actor if isinstance(actor, User) else None,
    )
    record_event(
        action="hub.reconciliation.source_uploaded",
        actor=actor,
        organization=organization,
        target=source,
        request=request,
        metadata={"kind": kind, "origin": origin, "size_bytes": len(content)},
    )
    return source, run, True


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ReconciliationError("Não foi possível identificar o texto do CSV.")


def _parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for pattern in DATE_FORMATS:
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    raise ReconciliationError(f"Data inválida: {text[:40]}")


def _parse_cents(value: object) -> int:
    text = str(value or "").strip().replace("R$", "").replace(" ", "")
    if not text:
        raise ReconciliationError("Valor ausente.")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return int((Decimal(text) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except InvalidOperation as exc:
        raise ReconciliationError(f"Valor inválido: {str(value)[:40]}") from exc


def _normalized_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def _auto_mapping(headers: list[str]) -> dict[str, str]:
    aliases = {
        "date": {"data", "date", "dt", "datamovimento", "datapostagem"},
        "amount": {"valor", "amount", "value", "valormovimento"},
        "debit": {"debito", "debitos", "saida", "saidas"},
        "credit": {"credito", "creditos", "entrada", "entradas"},
        "description": {"descricao", "historico", "memo", "description", "narrativa"},
        "document": {"documento", "doc", "numero", "numerodocumento", "referencia"},
        "counterparty": {"contraparte", "favorecido", "fornecedor", "cliente", "beneficiario"},
    }
    normalized = {_normalized_header(header): header for header in headers}
    result: dict[str, str] = {}
    for field, choices in aliases.items():
        for choice in choices:
            if choice in normalized:
                result[field] = normalized[choice]
                break
    return result


def _csv_dialect(text: str) -> type[csv.Dialect]:
    """Detect a dialect, with a deterministic fallback for short bank exports."""
    sample = text[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        first_line = next((line for line in text.splitlines() if line.strip()), "")
        delimiter = max((";", ",", "\t", "|"), key=first_line.count)

        class FallbackDialect(csv.excel):
            pass

        FallbackDialect.delimiter = delimiter if first_line.count(delimiter) else ";"
        return FallbackDialect


def _source_bytes(source: ReconciliationSourceFile) -> bytes:
    source.content.open("rb")
    try:
        source.content.seek(0)
        return cast(bytes, source.content.read())
    finally:
        source.content.close()


def _load_xlsx_workbook(content: bytes) -> Any:
    """Open a workbook without exposing archive/parser errors to an operator."""
    from openpyxl import load_workbook  # type: ignore[import-untyped]
    from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

    try:
        return load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise ReconciliationError(
            "Não foi possível abrir a planilha XLSX. Verifique se o arquivo não está corrompido."
        ) from exc


def preview_tabular(source: ReconciliationSourceFile) -> dict[str, Any]:
    content = _source_bytes(source)
    if source.kind == ReconciliationSourceFile.Kind.CSV:
        text = _decode_csv(content)
        dialect = _csv_dialect(text)
        reader = csv.reader(io.StringIO(text), dialect)
        rows = list(islice(reader, 51))
        headers = [str(item).strip() for item in rows[0]] if rows else []
        return {
            "headers": headers,
            "rows": rows[1:],
            "mapping": _auto_mapping(headers),
            "signature": _signature(headers),
        }
    if source.kind == ReconciliationSourceFile.Kind.XLSX:
        workbook = _load_xlsx_workbook(content)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True, max_row=51))
        headers = [str(item or "").strip() for item in rows[0]] if rows else []
        return {
            "headers": headers,
            "rows": [["" if item is None else str(item) for item in row] for row in rows[1:]],
            "mapping": _auto_mapping(headers),
            "signature": _signature(headers),
            "sheet": worksheet.title,
        }
    raise ReconciliationError("Prévia de mapeamento está disponível para CSV e XLSX.")


def _signature(headers: Iterable[str]) -> str:
    return hashlib.sha256(
        "\x1f".join(_normalized_header(item) for item in headers).encode()
    ).hexdigest()


def _compatible_layout_for_source(
    source: ReconciliationSourceFile,
) -> ReconciliationLayout | None:
    """Return the sole active layout with the same tabular structure.

    Layouts are a convenience, never an implicit decision between competing
    interpretations.  A malformed tabular file is still accepted as a source
    so the operator can inspect and map it, but it cannot inherit a layout.
    """
    if source.kind not in {
        ReconciliationSourceFile.Kind.CSV,
        ReconciliationSourceFile.Kind.XLSX,
    }:
        return None
    try:
        signature = preview_tabular(source)["signature"]
    except ReconciliationError:
        return None
    candidates = list(
        ReconciliationLayout.objects.filter(
            organization=source.organization,
            company=source.company,
            kind=source.kind,
            active=True,
            header_signature=signature,
        ).order_by("name", "-version", "-created_at")
    )
    return candidates[0] if len(candidates) == 1 else None


def save_layout(
    *,
    source: ReconciliationSourceFile,
    name: str,
    configuration: dict[str, Any],
    actor: object = None,
) -> ReconciliationLayout:
    preview = preview_tabular(source)
    _validate_tabular_mapping(configuration.get("mapping"), headers=preview["headers"])
    with transaction.atomic():
        previous = (
            ReconciliationLayout.objects.filter(
                organization=source.organization,
                company=source.company,
                kind=source.kind,
                name=name,
            )
            .order_by("-version")
            .first()
        )
        layout = ReconciliationLayout.objects.create(
            organization=source.organization,
            company=source.company,
            kind=source.kind,
            name=name[:120],
            version=(previous.version + 1) if previous else 1,
            header_signature=preview["signature"],
            configuration=configuration,
            created_by=actor if isinstance(actor, User) else None,
        )
        ReconciliationLayout.objects.filter(
            organization=source.organization,
            company=source.company,
            kind=source.kind,
            name=name[:120],
            active=True,
        ).exclude(id=layout.id).update(active=False)
    return layout


def _validate_tabular_mapping(mapping: object, *, headers: Iterable[str]) -> None:
    if not isinstance(mapping, dict):
        raise ReconciliationError("Informe o mapeamento das colunas antes de salvar o layout.")
    allowed_fields = {
        "date",
        "amount",
        "debit",
        "credit",
        "description",
        "document",
        "counterparty",
    }
    unexpected = set(mapping) - allowed_fields
    if unexpected:
        raise ReconciliationError("O mapeamento contém um campo que não é aceito.")
    selected = {
        field: value.strip()
        for field, value in mapping.items()
        if isinstance(value, str) and value.strip()
    }
    if any(not isinstance(value, str) for value in mapping.values()):
        raise ReconciliationError("Cada campo mapeado deve apontar para uma coluna de texto.")
    available_headers = {str(header) for header in headers}
    if any(header not in available_headers for header in selected.values()):
        raise ReconciliationError("O mapeamento contém uma coluna que não existe neste arquivo.")
    if len(set(selected.values())) != len(selected):
        raise ReconciliationError("Uma coluna não pode preencher mais de um campo normalizado.")
    if not selected.get("date") or not selected.get("description"):
        raise ReconciliationError("Mapeie ao menos data e histórico antes de salvar o layout.")
    if not any(selected.get(field) for field in ("amount", "debit", "credit")):
        raise ReconciliationError("Mapeie valor com sinal ou uma coluna de débito/crédito.")


@transaction.atomic
def prepare_run_for_layout(
    *, source: ReconciliationSourceFile, layout: ReconciliationLayout
) -> list[ReconciliationRun]:
    """Attach a validated layout to runs waiting specifically for mapping."""
    if layout.organization_id != source.organization_id or layout.company_id != source.company_id:
        raise ReconciliationError("O layout não pertence à empresa do arquivo.")
    if layout.kind != source.kind:
        raise ReconciliationError("O formato do layout não corresponde ao arquivo.")
    if layout.header_signature != preview_tabular(source)["signature"]:
        raise ReconciliationError("A estrutura do arquivo mudou; revise o mapeamento.")
    runs = list(
        ReconciliationRun.objects.select_for_update()
        .filter(source_file=source)
        .filter(
            Q(state__in=[ReconciliationRun.State.WAITING, ReconciliationRun.State.FAILED])
            | Q(state=ReconciliationRun.State.REVIEW, stage="mapping")
        )
        .order_by("created_at")
    )
    for run in runs:
        run.layout_version = layout
        run.state = ReconciliationRun.State.WAITING
        run.stage = "queued"
        run.errors = []
        run.save(update_fields=["layout_version", "state", "stage", "errors", "updated_at"])
    return runs


def _records_from_tabular(
    source: ReconciliationSourceFile, layout: ReconciliationLayout | None
) -> list[ExtractedRecord]:
    preview = preview_tabular(source)
    mapping = (layout.configuration.get("mapping") if layout else None) or preview["mapping"]
    if not isinstance(mapping, dict) or not mapping.get("date") or not mapping.get("description"):
        raise ReconciliationError("Mapeie ao menos data e histórico antes de processar o arquivo.")
    if layout and layout.header_signature != preview["signature"]:
        raise ReconciliationError("A estrutura do arquivo mudou; revise o mapeamento.")
    rows: Iterable[dict[str, Any]]
    content = _source_bytes(source)
    if source.kind == ReconciliationSourceFile.Kind.CSV:
        text = _decode_csv(content)
        dialect = _csv_dialect(text)
        rows = csv.DictReader(io.StringIO(text), dialect=dialect)
    else:
        workbook = _load_xlsx_workbook(content)
        sheet_name = str(
            (layout.configuration if layout else {}).get("sheet") or workbook.active.title
        )
        worksheet = workbook[sheet_name]
        iterator = worksheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        rows = (dict(zip(headers, row, strict=False)) for row in iterator)
    records: list[ExtractedRecord] = []
    for row_number, row in enumerate(rows, start=2):
        if row_number > MAX_ROWS + 1:
            raise ReconciliationError("Arquivo excede 100.000 linhas.")
        try:
            raw_date = row.get(str(mapping["date"]), "")
            amount_value = (
                row.get(str(mapping.get("amount", "")), "") if mapping.get("amount") else ""
            )
            if not amount_value:
                debit = (
                    _parse_cents(row.get(str(mapping.get("debit", "")), "0"))
                    if mapping.get("debit")
                    else 0
                )
                credit = (
                    _parse_cents(row.get(str(mapping.get("credit", "")), "0"))
                    if mapping.get("credit")
                    else 0
                )
                if debit and credit:
                    raise ReconciliationError("A linha tem débito e crédito simultâneos.")
                amount = credit or -debit
                raw_amount = str(row.get(str(mapping.get("credit" if credit else "debit", "")), ""))
            else:
                amount = _parse_cents(amount_value)
                raw_amount = str(amount_value)
            description = str(row.get(str(mapping["description"]), "")).strip()
            if not description:
                raise ReconciliationError("Histórico ausente.")
            records.append(
                ExtractedRecord(
                    source_key=f"row:{row_number}",
                    occurred_on=_parse_date(raw_date),
                    original_date=str(raw_date),
                    amount_cents=amount,
                    original_amount=raw_amount,
                    description=description,
                    document_number=str(row.get(str(mapping.get("document", "")), "")).strip(),
                    counterparty=str(row.get(str(mapping.get("counterparty", "")), "")).strip(),
                    source_reference={"row": row_number, "sheet": preview.get("sheet", "")},
                )
            )
        except ReconciliationError:
            raise
    return records


def _records_from_ofx(source: ReconciliationSourceFile) -> list[ExtractedRecord]:
    from apps.hub.reconciliation import parse_ofx

    parsed = parse_ofx(_source_bytes(source))
    return [
        ExtractedRecord(
            source_key=f"ofx:{item.external_id}",
            occurred_on=item.occurred_on,
            original_date=item.occurred_on.isoformat(),
            amount_cents=item.amount_cents,
            original_amount=str(Decimal(item.amount_cents) / 100),
            description=item.description,
            source_reference={"fitid": item.external_id, "account": parsed.account_reference},
        )
        for item in parsed.transactions
    ]


def _records_from_pdf(source: ReconciliationSourceFile) -> list[ExtractedRecord]:
    content = _source_bytes(source)
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            if len(pdf.pages) > MAX_PAGES:
                raise ReconciliationError("PDF excede 500 páginas.")
            records: list[ExtractedRecord] = []
            line_number = 0
            date_pattern = re.compile(r"(\d{2}[/-]\d{2}[/-]\d{2,4})")
            value_pattern = re.compile(
                r"(?<!\d)(-?\s*\d{1,3}(?:\.\d{3})*,\d{2}|-?\s*\d+\.\d{2})(?!\d)"
            )
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                words = page.extract_words() or []
                quality = 100 if text else 0
                for raw_line in text.splitlines():
                    line_number += 1
                    date_match = date_pattern.search(raw_line)
                    amount_match = value_pattern.search(raw_line)
                    if not date_match or not amount_match:
                        continue
                    try:
                        occurred_on = _parse_date(date_match.group(1))
                        amount_cents = _parse_cents(amount_match.group(1))
                    except ReconciliationError:
                        continue
                    description = (
                        raw_line[date_match.end() : amount_match.start()].strip(" -")
                        or raw_line[:180]
                    )
                    matched_words = [
                        word
                        for word in words
                        if str(word.get("text", "")).strip()
                        in {date_match.group(1), amount_match.group(1).strip()}
                    ]
                    source_reference: dict[str, Any] = {
                        "page": page_number,
                        "line": line_number,
                        "word_count": len(words),
                    }
                    if matched_words:
                        source_reference["bbox"] = {
                            "x0": min(float(word["x0"]) for word in matched_words),
                            "top": min(float(word["top"]) for word in matched_words),
                            "x1": max(float(word["x1"]) for word in matched_words),
                            "bottom": max(float(word["bottom"]) for word in matched_words),
                            "unit": "pdf_points",
                        }
                    records.append(
                        ExtractedRecord(
                            source_key=f"pdf:{page_number}:{line_number}",
                            occurred_on=occurred_on,
                            original_date=date_match.group(1),
                            amount_cents=amount_cents,
                            original_amount=amount_match.group(1),
                            description=description[:1000],
                            source_reference=source_reference,
                            confidence=quality,
                            confidence_details={"method": "native_pdf_text", "page": page_number},
                        )
                    )
            if records:
                return records
    except ImportError:
        pass
    return _records_from_scanned_pdf(content)


def _records_from_scanned_pdf(content: bytes) -> list[ExtractedRecord]:
    """Use a locally installed Tesseract only; unavailable OCR remains a review state."""
    tesseract_path = _tesseract_executable()
    if tesseract_path is None or not local_ocr_available():
        return []
    try:
        import pypdfium2  # type: ignore[import-untyped]

        document = pypdfium2.PdfDocument(content)
        if len(document) > MAX_PAGES:
            raise ReconciliationError("PDF excede 500 páginas.")
        records: list[ExtractedRecord] = []
        date_pattern = re.compile(r"(\d{2}[/-]\d{2}[/-]\d{2,4})")
        value_pattern = re.compile(r"(?<!\d)(-?\s*\d{1,3}(?:\.\d{3})*,\d{2}|-?\s*\d+\.\d{2})(?!\d)")
        with tempfile.TemporaryDirectory(prefix="cica-reconciliation-ocr-") as temporary_directory:
            for page_number in range(1, len(document) + 1):
                image_path = PurePath(temporary_directory, f"page-{page_number}.png")
                page = document[page_number - 1]
                page.render(scale=2).to_pil().save(image_path)
                result = subprocess.run(  # noqa: S603 - fixed local binary and private temp file
                    [
                        tesseract_path,
                        str(image_path),
                        "stdout",
                        "-l",
                        "por",
                        "--psm",
                        "6",
                        "-c",
                        "tessedit_create_tsv=1",
                    ],
                    capture_output=True,
                    check=False,
                    text=True,
                    timeout=90,
                    env=_tesseract_environment(),
                )
                if result.returncode:
                    continue
                # TSV columns 2-4 identify block, paragraph and line. Column 5 is
                # the word ordinal, so including it would incorrectly turn every
                # word into an independent OCR "line".
                words: dict[tuple[str, str, str], list[tuple[str, int, int, int, int, int]]] = {}
                for row in result.stdout.splitlines()[1:]:
                    parts = row.split("\t")
                    if len(parts) != 12 or not parts[11].strip():
                        continue
                    try:
                        confidence = max(0, min(100, int(float(parts[10]))))
                        left, top, width, height = (int(value) for value in parts[6:10])
                    except ValueError:
                        continue
                    line_key = (parts[2], parts[3], parts[4])
                    words.setdefault(line_key, []).append(
                        (parts[11], confidence, left, top, width, height)
                    )
                for line_number, values in enumerate(words.values(), start=1):
                    raw_line = " ".join(word for word, *_ in values)
                    date_match = date_pattern.search(raw_line)
                    amount_match = value_pattern.search(raw_line)
                    if not date_match or not amount_match:
                        continue
                    try:
                        occurred_on = _parse_date(date_match.group(1))
                        amount_cents = _parse_cents(amount_match.group(1))
                    except ReconciliationError:
                        continue
                    description = (
                        raw_line[date_match.end() : amount_match.start()].strip(" -")
                        or raw_line[:180]
                    )
                    confidence = round(sum(value[1] for value in values) / len(values))
                    source_reference = {
                        "page": page_number,
                        "ocr_line": line_number,
                        "bbox": {
                            "x0": min(value[2] for value in values),
                            "top": min(value[3] for value in values),
                            "x1": max(value[2] + value[4] for value in values),
                            "bottom": max(value[3] + value[5] for value in values),
                            "unit": "pixels@2x",
                        },
                    }
                    records.append(
                        ExtractedRecord(
                            source_key=f"ocr:{page_number}:{line_number}",
                            occurred_on=occurred_on,
                            original_date=date_match.group(1),
                            amount_cents=amount_cents,
                            original_amount=amount_match.group(1),
                            description=description[:1000],
                            source_reference=source_reference,
                            confidence=confidence,
                            confidence_details={
                                "method": "tesseract-local",
                                "page": page_number,
                                "reading_quality": confidence,
                            },
                        )
                    )
        return records
    except (ImportError, OSError, subprocess.TimeoutExpired):
        return []


def _record_direction(amount_cents: int) -> str:
    return (
        NormalizedMovement.Direction.INFLOW
        if amount_cents >= 0
        else NormalizedMovement.Direction.OUTFLOW
    )


def _matches_condition(condition: dict[str, Any], movement: NormalizedMovement) -> bool:
    field = str(condition.get("field", ""))
    operator = str(condition.get("operator", ""))
    expected = condition.get("value", "")
    financial_account = movement.financial_account
    values = {
        "financial_account": (
            financial_account.account_reference
            if financial_account is not None
            else ""
        ),
        "description": movement.description,
        "counterparty": movement.counterparty,
        "document": movement.document_number,
        "direction": movement.direction,
        "amount_cents": movement.amount_cents,
    }
    actual = values.get(field)
    if operator == "equals":
        return str(actual).casefold() == str(expected).casefold()
    if operator == "contains":
        return str(expected).casefold() in str(actual).casefold()
    if operator == "range" and isinstance(expected, dict):
        minimum = expected.get("min", -(10**18))
        maximum = expected.get("max", 10**18)
        if (
            not isinstance(minimum, (int, str))
            or not isinstance(actual, (int, str))
            or not isinstance(maximum, (int, str))
        ):
            return False
        try:
            return int(minimum) <= int(actual) <= int(maximum)
        except ValueError:
            return False
    if operator == "regex":
        pattern = str(expected)
        if len(pattern) > 160 or re.search(r"\(\?[:=!<].*[+*].*\)[+*]", pattern):
            return False
        try:
            return re.search(pattern, str(actual or ""), flags=re.IGNORECASE) is not None
        except re.error:
            return False
    return False


def apply_rules(movement: NormalizedMovement, *, actor: object = None) -> NormalizedMovement:
    rules = ReconciliationRule.objects.filter(
        organization=movement.organization,
        company=movement.company,
        state=ReconciliationRule.State.ACTIVE,
    ).order_by("priority", "created_at", "id")
    matching: list[ReconciliationRule] = []
    for rule in rules:
        all_conditions = [item for item in rule.all_conditions if isinstance(item, dict)]
        any_conditions = [item for item in rule.any_conditions if isinstance(item, dict)]
        if all(_matches_condition(item, movement) for item in all_conditions) and (
            not any_conditions or any(_matches_condition(item, movement) for item in any_conditions)
        ):
            matching.append(rule)
    if not matching:
        return movement
    actions = [rule.actions for rule in matching]
    keys = (
        "debit_account_code",
        "credit_account_code",
        "cost_center_code",
        "accounting_history",
        "review",
        "ignore",
    )
    conflicts = {
        key
        for key in keys
        if len({str(action.get(key, "")) for action in actions if action.get(key, "")}) > 1
    }
    if any(action.get("ignore") for action in actions) and any(
        not action.get("ignore") for action in actions
    ):
        conflicts.add("ignore")
    if conflicts:
        movement.review_state = NormalizedMovement.ReviewState.CONFLICT
        movement.confidence_details = {
            **movement.confidence_details,
            "rule_conflicts": sorted(conflicts),
            "rules": [str(rule.id) for rule in matching],
        }
        movement.save(update_fields=["review_state", "confidence_details", "updated_at"])
        return movement
    action = actions[0]
    movement.debit_account_code = str(
        action.get("debit_account_code", movement.debit_account_code)
    )[:64]
    movement.credit_account_code = str(
        action.get("credit_account_code", movement.credit_account_code)
    )[:64]
    movement.cost_center_code = str(action.get("cost_center_code", movement.cost_center_code))[:64]
    movement.accounting_history = str(
        action.get("accounting_history", movement.accounting_history)
    )[:500]
    movement.applied_rule = matching[0]
    movement.classification_source = NormalizedMovement.ClassificationSource.RULE
    movement.review_state = (
        NormalizedMovement.ReviewState.IGNORED
        if action.get("ignore")
        else NormalizedMovement.ReviewState.PENDING
        if action.get("review") or action.get("require_review")
        else NormalizedMovement.ReviewState.READY
    )
    movement.save()
    return movement


def require_review_for_low_quality_extraction(movement: NormalizedMovement) -> NormalizedMovement:
    """Never let a weak OCR reading become an unattended accounting decision."""

    if (
        movement.confidence_details.get("method") != "tesseract-local"
        or movement.confidence >= OCR_REVIEW_THRESHOLD
    ):
        return movement
    movement.review_state = NormalizedMovement.ReviewState.PENDING
    movement.confidence_details = {
        **movement.confidence_details,
        "requires_review": "low_ocr_quality",
        "review_threshold": OCR_REVIEW_THRESHOLD,
    }
    movement.save(update_fields=["review_state", "confidence_details", "updated_at"])
    return movement


def _match_token(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def auto_reconcile_unique_movement(movement: NormalizedMovement) -> MovementReconciliation | None:
    """Confirm only an exact, independent 1:1 match with documentary evidence."""
    if (
        movement.review_state != NormalizedMovement.ReviewState.READY
        or not movement.occurred_on
        or not movement.document_number.strip()
        or movement.amount_cents == 0
    ):
        return None
    document_token = _match_token(movement.document_number)
    if not document_token:
        return None
    candidates: list[JournalEntry] = []
    for entry in (
        JournalEntry.objects.filter(
            organization=movement.organization,
            company=movement.company,
            occurred_on=movement.occurred_on,
            state=JournalEntry.State.APPROVED,
        )
        .exclude(movement=movement)
        .prefetch_related("lines")
    ):
        if document_token not in _match_token(entry.history):
            continue
        if _entry_capacity(entry) != abs(movement.amount_cents):
            continue
        if MovementReconciliation.objects.filter(
            entry=entry, state=MovementReconciliation.State.CONFIRMED
        ).exists():
            continue
        candidates.append(entry)
    competing_movements = (
        NormalizedMovement.objects.filter(
            organization=movement.organization,
            company=movement.company,
            occurred_on=movement.occurred_on,
            amount_cents=movement.amount_cents,
        )
        .exclude(id=movement.id)
        .exclude(review_state=NormalizedMovement.ReviewState.IGNORED)
    )
    if len(candidates) != 1 or competing_movements.exists():
        return None
    reconciliation = confirm_reconciliation(
        movement=movement,
        entry=candidates[0],
        amount_cents=abs(movement.amount_cents),
        evidence={
            "method": "deterministic_exact_1to1",
            "document": movement.document_number,
            "date": movement.occurred_on.isoformat(),
        },
    )
    record_event(
        action="hub.reconciliation.auto_confirmed",
        organization=movement.organization,
        target=reconciliation,
        metadata={"movement_id": str(movement.id), "entry_id": str(candidates[0].id)},
    )
    return reconciliation


@transaction.atomic
def save_rule_from_movement(
    *, movement: NormalizedMovement, actor: User | None = None
) -> ReconciliationRule:
    """Turn an explicit correction into a narrow, auditable company rule."""
    movement = NormalizedMovement.objects.select_for_update().get(pk=movement.pk)
    if movement.review_state == NormalizedMovement.ReviewState.IGNORED:
        raise ReconciliationError("Desfaça o ignorar antes de transformar o movimento em regra.")
    if not movement.description.strip():
        raise ReconciliationError("O movimento precisa de histórico para virar uma regra.")
    if not movement.debit_account_code or not movement.credit_account_code:
        raise ReconciliationError("Informe as duas contas antes de salvar uma regra.")
    conditions = [
        {"field": "description", "operator": "equals", "value": movement.description},
        {"field": "direction", "operator": "equals", "value": movement.direction},
    ]
    financial_account = movement.financial_account
    if financial_account is not None and financial_account.account_reference:
        conditions.insert(
            0,
            {
                "field": "financial_account",
                "operator": "equals",
                "value": financial_account.account_reference,
            },
        )
    rule = ReconciliationRule.objects.create(
        organization=movement.organization,
        company=movement.company,
        name=f"{movement.description[:100]} ({movement.direction})",
        priority=100,
        state=ReconciliationRule.State.ACTIVE,
        all_conditions=conditions,
        actions={
            "debit_account_code": movement.debit_account_code,
            "credit_account_code": movement.credit_account_code,
            "cost_center_code": movement.cost_center_code,
            "accounting_history": movement.accounting_history,
        },
        created_by=actor,
    )
    movement.applied_rule = rule
    movement.save(update_fields=["applied_rule", "updated_at"])
    return rule


def _entry_capacity(entry: JournalEntry) -> int:
    debit = sum(line.amount_cents for line in entry.lines.filter(side=JournalLine.Side.DEBIT))
    credit = sum(line.amount_cents for line in entry.lines.filter(side=JournalLine.Side.CREDIT))
    if debit <= 0 or debit != credit:
        raise ReconciliationError("O lançamento de comparação não está equilibrado.")
    return debit


@transaction.atomic
def confirm_reconciliation(
    *,
    movement: NormalizedMovement,
    entry: JournalEntry,
    amount_cents: int,
    evidence: dict[str, Any],
    actor: User | None = None,
) -> MovementReconciliation:
    """Confirm an evidence-backed allocation without consuming either side twice."""
    movement = NormalizedMovement.objects.select_for_update().get(pk=movement.pk)
    entry = JournalEntry.objects.select_for_update().get(pk=entry.pk)
    if movement.organization_id != entry.organization_id or movement.company_id != entry.company_id:
        raise ReconciliationError("Movimento e lançamento precisam pertencer à mesma empresa.")
    if entry.movement_id == movement.id:
        raise ReconciliationError(
            "O lançamento gerado pelo próprio movimento não é evidência independente."
        )
    if amount_cents <= 0:
        raise ReconciliationError("Informe uma alocação positiva.")
    if not evidence:
        raise ReconciliationError("Informe a evidência que justifica a conciliação.")
    entry_capacity = _entry_capacity(entry)
    movement_used = sum(
        MovementReconciliation.objects.select_for_update()
        .filter(movement=movement, state=MovementReconciliation.State.CONFIRMED)
        .values_list("amount_cents", flat=True)
    )
    entry_used = sum(
        MovementReconciliation.objects.select_for_update()
        .filter(entry=entry, state=MovementReconciliation.State.CONFIRMED)
        .values_list("amount_cents", flat=True)
    )
    if movement_used + amount_cents > abs(movement.amount_cents):
        raise ReconciliationError("A alocação excede o valor disponível do movimento.")
    if entry_used + amount_cents > entry_capacity:
        raise ReconciliationError("A alocação excede o valor disponível do lançamento.")
    reconciliation, created = MovementReconciliation.objects.get_or_create(
        organization=movement.organization,
        movement=movement,
        entry=entry,
        defaults={
            "amount_cents": amount_cents,
            "state": MovementReconciliation.State.CONFIRMED,
            "evidence": evidence,
            "confirmed_at": timezone.now(),
            "confirmed_by": actor,
        },
    )
    if not created:
        raise ReconciliationError("Essa relação já existe; desfaça-a antes de alterar a alocação.")
    return reconciliation


@transaction.atomic
def undo_reconciliation(
    *, reconciliation: MovementReconciliation, actor: User | None = None
) -> None:
    locked = MovementReconciliation.objects.select_for_update().get(pk=reconciliation.pk)
    if locked.state != MovementReconciliation.State.CONFIRMED:
        return
    locked.state = MovementReconciliation.State.UNDONE
    locked.confirmed_at = None
    locked.confirmed_by = actor
    locked.save(update_fields=["state", "confirmed_at", "confirmed_by", "updated_at"])


def process_run(run_id: str) -> dict[str, int | str]:
    try:
        parsed_id = uuid.UUID(str(run_id))
    except ValueError:
        return {"state": "invalid_id"}
    token = uuid.uuid4()
    now = timezone.now()
    acquired = (
        ReconciliationRun.objects.filter(
            id=parsed_id,
            state__in=[ReconciliationRun.State.WAITING, ReconciliationRun.State.FAILED],
        )
        .filter(models_Q_lease(now))
        .update(
            state=ReconciliationRun.State.PROCESSING,
            stage="extracting",
            lease_token=token,
            lease_until=now + timedelta(minutes=20),
            started_at=now,
        )
    )
    if not acquired:
        return {"state": "not_ready"}
    try:
        run = ReconciliationRun.objects.select_related(
            "source_file", "source_file__company", "source_file__financial_account"
        ).get(id=parsed_id)
        source = run.source_file
        if run.cancel_requested_at:
            _finish_run(run, token, ReconciliationRun.State.CANCELED, "canceled")
            return {"state": "canceled"}
        if run.checkpoint.get("mode") == "rules":
            return _reapply_rules_for_source(run, token)
        if source.kind == ReconciliationSourceFile.Kind.OFX:
            records = _records_from_ofx(source)
            financial_account = source.financial_account
            if financial_account is not None:
                reported_accounts = {
                    str(record.source_reference.get("account", "")).strip()
                    for record in records
                    if record.source_reference
                }
                selected_account = financial_account.account_reference.strip()
                if reported_accounts != {selected_account}:
                    _finish_run(
                        run,
                        token,
                        ReconciliationRun.State.REVIEW,
                        "financial_account_review",
                        errors=[
                            {
                                "code": "financial_account_mismatch",
                                "message": (
                                    "A conta informada pelo OFX não corresponde à conta financeira "
                                    "selecionada. Revise a associação antes de normalizar "
                                    "o arquivo."
                                ),
                            }
                        ],
                    )
                    return {"state": "review", "created": 0}
        elif source.kind in {ReconciliationSourceFile.Kind.CSV, ReconciliationSourceFile.Kind.XLSX}:
            if run.layout_version_id is None:
                _finish_run(
                    run,
                    token,
                    ReconciliationRun.State.REVIEW,
                    "mapping",
                    errors=[
                        {
                            "code": "mapping_required",
                            "message": (
                                "Confirme o mapeamento das colunas antes de "
                                "normalizar este arquivo."
                            ),
                        }
                    ],
                )
                return {"state": "review", "created": 0}
            records = _records_from_tabular(source, run.layout_version)
        else:
            records = _records_from_pdf(source)
            if not records:
                _finish_run(
                    run,
                    token,
                    ReconciliationRun.State.REVIEW,
                    "ocr_review",
                    errors=[
                        {
                            "code": "ocr_required",
                            "message": (
                                "PDF sem texto extraível; configure OCR local "
                                "ou revise o documento."
                            ),
                        }
                    ],
                )
                return {"state": "review", "created": 0}
        created = 0
        ignored = 0
        errors: list[dict[str, str]] = []
        auto_reconciliation_candidates: list[uuid.UUID] = []
        for index, record in enumerate(records, start=1):
            if ReconciliationRun.objects.filter(
                id=run.id, lease_token=token, cancel_requested_at__isnull=False
            ).exists():
                _finish_run(
                    run,
                    token,
                    ReconciliationRun.State.CANCELED,
                    "canceled",
                    processed=index - 1,
                    created=created,
                    ignored=ignored,
                    errors=errors,
                )
                return {"state": "canceled", "created": created}
            try:
                movement, was_created = NormalizedMovement.objects.get_or_create(
                    organization=run.organization,
                    source_file=source,
                    source_key=record.source_key,
                    defaults={
                        "run": run,
                        "company": source.company,
                        "financial_account": source.financial_account,
                        "source_reference": record.source_reference or {},
                        "original_date": record.original_date,
                        "occurred_on": record.occurred_on,
                        "original_amount": record.original_amount,
                        "amount_cents": record.amount_cents,
                        "direction": _record_direction(record.amount_cents),
                        "original_description": record.description,
                        "description": record.description,
                        "document_number": record.document_number,
                        "counterparty": record.counterparty,
                        "confidence": record.confidence,
                        "confidence_details": record.confidence_details or {},
                    },
                )
                if was_created:
                    created += 1
                    apply_rules(movement)
                    require_review_for_low_quality_extraction(movement)
                    # Evaluate automatic reconciliation only after every record in
                    # this source has been persisted.  A duplicate later in the
                    # same input must make the candidate ambiguous.
                    if movement.review_state == NormalizedMovement.ReviewState.READY:
                        auto_reconciliation_candidates.append(movement.id)
                    if movement.review_state == NormalizedMovement.ReviewState.IGNORED:
                        ignored += 1
            except (IntegrityError, ReconciliationError) as exc:
                errors.append(
                    {"line": str(index), "code": "record_invalid", "message": str(exc)[:240]}
                )
        for movement_id in auto_reconciliation_candidates:
            try:
                auto_reconcile_unique_movement(NormalizedMovement.objects.get(id=movement_id))
            except ReconciliationError:
                # A concurrent review may have consumed the candidate after this
                # run's deterministic check. Leave it pending for manual review.
                continue
        state = (
            ReconciliationRun.State.REVIEW
            if created or errors
            else ReconciliationRun.State.COMPLETED
        )
        if errors and created == 0:
            state = ReconciliationRun.State.FAILED
        elif errors:
            state = ReconciliationRun.State.COMPLETED_ALERTS
        _finish_run(
            run,
            token,
            state,
            "review" if state == ReconciliationRun.State.REVIEW else "completed",
            processed=len(records),
            created=created,
            ignored=ignored,
            errors=errors,
        )
        return {"state": state, "created": created, "errors": len(errors)}
    except Exception as exc:
        failed_run = ReconciliationRun.objects.filter(id=parsed_id).first()
        if failed_run:
            _finish_run(
                failed_run,
                token,
                ReconciliationRun.State.FAILED,
                "failed",
                errors=[{"code": "processing_failed", "message": str(exc)[:240]}],
            )
        return {"state": "failed"}


def _reapply_rules_for_source(run: ReconciliationRun, token: uuid.UUID) -> dict[str, int | str]:
    """Apply current rules only where no protected accounting decision exists."""

    source = run.source_file
    all_movements = NormalizedMovement.objects.filter(source_file=source)
    eligible = (
        all_movements.exclude(classification_source=NormalizedMovement.ClassificationSource.MANUAL)
        .exclude(
            journal_entries__state__in=[
                JournalEntry.State.DRAFT,
                JournalEntry.State.APPROVED,
                JournalEntry.State.EXPORTED,
            ]
        )
        .exclude(reconciliations__state=MovementReconciliation.State.CONFIRMED)
        .order_by("id")
        .distinct()
    )
    total = all_movements.count()
    eligible_count = eligible.count()
    protected_count = total - eligible_count
    changed = 0
    for index, movement_id in enumerate(
        eligible.values_list("id", flat=True).iterator(100), start=1
    ):
        if ReconciliationRun.objects.filter(
            id=run.id, lease_token=token, cancel_requested_at__isnull=False
        ).exists():
            _finish_run(
                run,
                token,
                ReconciliationRun.State.CANCELED,
                "canceled",
                processed=index - 1,
                updated=changed,
            )
            return {"state": "canceled", "updated": changed}
        with transaction.atomic():
            movement = NormalizedMovement.objects.select_for_update().get(id=movement_id)
            before = (
                movement.debit_account_code,
                movement.credit_account_code,
                movement.cost_center_code,
                movement.accounting_history,
                movement.review_state,
                movement.classification_source,
                movement.applied_rule_id,
                movement.confidence_details,
            )
            details = {
                key: value
                for key, value in movement.confidence_details.items()
                if key not in {"rule_conflicts", "rules"}
            }
            movement.debit_account_code = ""
            movement.credit_account_code = ""
            movement.cost_center_code = ""
            movement.accounting_history = ""
            movement.applied_rule = None
            movement.classification_source = NormalizedMovement.ClassificationSource.NONE
            movement.review_state = NormalizedMovement.ReviewState.PENDING
            movement.confidence_details = details
            movement.save()
            apply_rules(movement)
            require_review_for_low_quality_extraction(movement)
            after = (
                movement.debit_account_code,
                movement.credit_account_code,
                movement.cost_center_code,
                movement.accounting_history,
                movement.review_state,
                movement.classification_source,
                movement.applied_rule_id,
                movement.confidence_details,
            )
            if after != before:
                movement.revision += 1
                movement.save(update_fields=["revision", "updated_at"])
                changed += 1
    awaiting_review = eligible.filter(
        review_state__in=[
            NormalizedMovement.ReviewState.PENDING,
            NormalizedMovement.ReviewState.CONFLICT,
        ]
    ).exists()
    errors = (
        [
            {
                "code": "protected_movements",
                "message": (
                    f"{protected_count} movimento(s) com revisão manual, conciliação ou "
                    "lançamento foram preservados."
                ),
            }
        ]
        if protected_count
        else []
    )
    state = (
        ReconciliationRun.State.REVIEW
        if awaiting_review
        else ReconciliationRun.State.COMPLETED_ALERTS
        if errors
        else ReconciliationRun.State.COMPLETED
    )
    _finish_run(
        run,
        token,
        state,
        "rules_reapplied",
        processed=eligible_count,
        updated=changed,
        errors=errors,
    )
    return {"state": state, "updated": changed, "protected": protected_count}


def models_Q_lease(now: Any) -> Any:
    from django.db.models import Q

    return Q(lease_until__isnull=True) | Q(lease_until__lt=now)


def _finish_run(
    run: ReconciliationRun,
    token: uuid.UUID,
    state: str,
    stage: str,
    *,
    processed: int = 0,
    created: int = 0,
    updated: int = 0,
    ignored: int = 0,
    errors: list[dict[str, str]] | None = None,
) -> None:
    errors = errors or []
    ReconciliationRun.objects.filter(id=run.id, lease_token=token).update(
        state=state,
        stage=stage,
        processed_count=processed,
        total_count=max(processed, run.total_count),
        created_count=created,
        updated_count=updated,
        ignored_count=ignored,
        error_count=len(errors),
        errors=errors,
        completed_at=timezone.now(),
        lease_token=None,
        lease_until=None,
    )


def validate_journal_entry(entry: JournalEntry) -> list[str]:
    errors: list[str] = []
    if entry.movement_id:
        movement = entry.movement
        if entry.source_movement_revision is None:
            errors.append("O lançamento não possui a revisão de origem para validação.")
        elif movement is None or movement.revision != entry.source_movement_revision:
            errors.append("O movimento de origem mudou; gere e aprove um novo lançamento.")
    if not entry.history.strip():
        errors.append("Informe o histórico contábil.")
    lines = list(entry.lines.all())
    if len(lines) < 2:
        errors.append("O lançamento precisa de ao menos duas partidas.")
    debit = sum(line.amount_cents for line in lines if line.side == JournalLine.Side.DEBIT)
    credit = sum(line.amount_cents for line in lines if line.side == JournalLine.Side.CREDIT)
    if debit != credit or debit <= 0:
        errors.append("Débitos e créditos devem ser iguais e maiores que zero.")
    valid_accounts = set(
        LedgerAccount.objects.filter(
            organization=entry.organization,
            company=entry.company,
            active=True,
            accepts_entries=True,
        ).values_list("code", flat=True)
    )
    for line in lines:
        if line.account_code not in valid_accounts:
            errors.append(f"Conta inválida ou inativa: {line.account_code}.")
    if AccountingPeriod.objects.filter(
        organization=entry.organization,
        company=entry.company,
        starts_on__lte=entry.occurred_on,
        ends_on__gte=entry.occurred_on,
        locked_at__isnull=False,
    ).exists():
        errors.append("O período contábil está bloqueado.")
    return errors


@transaction.atomic
def create_journal_entry_from_movement(
    *, movement: NormalizedMovement, actor: User | None = None
) -> JournalEntry:
    """Create one balanced draft from an explicitly classified movement."""
    movement = NormalizedMovement.objects.select_for_update().get(pk=movement.pk)
    if movement.review_state == NormalizedMovement.ReviewState.IGNORED:
        raise ReconciliationError("Desfaça o ignorar antes de gerar um lançamento.")
    if (
        not movement.occurred_on
        or not movement.debit_account_code
        or not movement.credit_account_code
    ):
        raise ReconciliationError(
            "Informe data, contas de débito e crédito antes de gerar o lançamento."
        )
    if movement.amount_cents == 0:
        raise ReconciliationError("Um movimento de valor zero não gera lançamento.")
    existing = movement.journal_entries.exclude(state=JournalEntry.State.INVALID).first()
    if existing:
        return existing
    entry = JournalEntry.objects.create(
        organization=movement.organization,
        company=movement.company,
        movement=movement,
        source_movement_revision=movement.revision,
        occurred_on=movement.occurred_on,
        history=movement.accounting_history or movement.description,
        purpose="record",
    )
    amount_cents = abs(movement.amount_cents)
    JournalLine.objects.bulk_create(
        [
            JournalLine(
                organization=movement.organization,
                entry=entry,
                account_code=movement.debit_account_code,
                side=JournalLine.Side.DEBIT,
                amount_cents=amount_cents,
                cost_center_code=movement.cost_center_code,
            ),
            JournalLine(
                organization=movement.organization,
                entry=entry,
                account_code=movement.credit_account_code,
                side=JournalLine.Side.CREDIT,
                amount_cents=amount_cents,
                cost_center_code=movement.cost_center_code,
            ),
        ]
    )
    return entry


def approve_journal_entry(
    *, entry: JournalEntry, actor: object = None, request: object = None
) -> JournalEntry:
    """Approve a validated snapshot and persist an invalid state on failure."""

    failure = ""
    with transaction.atomic():
        locked = (
            # ``movement`` is optional.  PostgreSQL rejects a blanket ``FOR UPDATE``
            # when ``select_related`` introduces that outer join, so lock only the
            # journal entry whose approval is being decided.
            JournalEntry.objects.select_for_update(of=("self",))
            .select_related("company", "movement")
            .get(id=entry.id)
        )
        errors = validate_journal_entry(locked)
        if errors:
            locked.state = JournalEntry.State.INVALID
            locked.save(update_fields=["state", "updated_at"])
            failure = " ".join(errors)
        else:
            locked.state = JournalEntry.State.APPROVED
            locked.approved_at = timezone.now()
            locked.approved_by = actor if isinstance(actor, User) else None
            locked.save(update_fields=["state", "approved_at", "approved_by", "updated_at"])
            record_event(
                action="hub.reconciliation.journal_approved",
                actor=actor,
                organization=locked.organization,
                target=locked,
                request=request,
                metadata={"line_count": locked.lines.count()},
            )
    if failure:
        raise ReconciliationError(failure)
    return locked


def render_dominio_csv(entries: Iterable[JournalEntry]) -> bytes:
    """Produces the documented ten-column model; importer confirmation remains pending."""
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(
        [
            "Data",
            "Cód. Conta Debito",
            "Cód. Conta Credito",
            "Valor",
            "Cód. Histórico",
            "Complemento Histórico",
            "Inicia Lote",
            "Código Matriz/Filial",
            "Centro de Custo Débito",
            "Centro de Custo Crédito",
        ]
    )
    for entry in entries:
        debits = [line for line in entry.lines.all() if line.side == JournalLine.Side.DEBIT]
        credits = [line for line in entry.lines.all() if line.side == JournalLine.Side.CREDIT]
        # The official model accepts simple and multiple entries; expand an entry by
        # pairing its debit and credit parts while preserving their total amounts.
        if len(debits) != len(credits):
            raise ReconciliationError(
                "Partidas compostas com quantidades diferentes exigem agrupamento homologado "
                "para o Domínio."
            )
        for index, (debit, credit) in enumerate(zip(debits, credits, strict=False)):
            if debit.amount_cents != credit.amount_cents:
                raise ReconciliationError(
                    "Partidas múltiplas exigem um agrupamento homologado para o Domínio."
                )
            writer.writerow(
                [
                    entry.occurred_on.strftime("%d/%m/%Y"),
                    debit.account_code,
                    credit.account_code,
                    f"{Decimal(debit.amount_cents) / 100:.2f}",
                    "",
                    entry.history,
                    "1" if index == 0 else "",
                    "",
                    debit.cost_center_code,
                    credit.cost_center_code,
                ]
            )
    return output.getvalue().encode("utf-8-sig")


@dataclass(frozen=True)
class AccountingExportAdapter:
    """A reviewed target contract; no target falls back to another system."""

    target: str
    version: str
    render: Callable[[Iterable[JournalEntry]], bytes]


ACCOUNTING_EXPORT_ADAPTERS: dict[str, AccountingExportAdapter] = {
    AccountingExport.Target.DOMINIO: AccountingExportAdapter(
        target=AccountingExport.Target.DOMINIO,
        version="dominio-3.1-pending-homologation",
        render=render_dominio_csv,
    ),
}


def get_accounting_export_adapter(target: str) -> AccountingExportAdapter:
    """Return only an explicitly registered, reviewed export target."""
    adapter = ACCOUNTING_EXPORT_ADAPTERS.get(target)
    if adapter is None:
        if target == AccountingExport.Target.SIESCON:
            raise ReconciliationError(
                "Exportação Siescon está bloqueada até registrar layout e adaptador revisados."
            )
        raise ReconciliationError("Destino de exportação contábil inválido.")
    return adapter


@transaction.atomic
def create_export(
    *,
    organization: Organization,
    company: Any,
    start: date,
    end: date,
    actor: object = None,
    reexport_of: AccountingExport | None = None,
    reexport_reason: str = "",
    target: str = AccountingExport.Target.DOMINIO,
    request: object = None,
) -> AccountingExport:
    adapter = get_accounting_export_adapter(target)
    if (
        target == AccountingExport.Target.DOMINIO
        and not settings.RECONCILIATION_DOMINIO_EXPORT_HOMOLOGATED
    ):
        raise ReconciliationError(
            "Exportação Domínio está bloqueada até a homologação real do adaptador."
        )
    query = (
        # ``movement`` is optional, therefore this join is outer.  Locking the
        # entry alone keeps export serialization valid on PostgreSQL.
        JournalEntry.objects.select_for_update(of=("self",))
        .select_related("movement")
        .filter(
            organization=organization,
            company=company,
            occurred_on__range=(start, end),
        )
    )
    if reexport_of is None:
        query = query.filter(state=JournalEntry.State.APPROVED)
    else:
        if (
            reexport_of.organization_id != organization.id
            or reexport_of.company_id != company.id
            or not reexport_reason.strip()
        ):
            raise ReconciliationError("Reexportação exige exportação original da empresa e motivo.")
        query = query.filter(id__in=reexport_of.entry_ids, state=JournalEntry.State.EXPORTED)
    entries = list(query.prefetch_related("lines").order_by("occurred_on", "id"))
    if not entries:
        raise ReconciliationError("Não há lançamentos aprovados novos para exportar.")
    if (
        reexport_of is None
        and JournalEntry.objects.filter(
            id__in=[entry.id for entry in entries], state=JournalEntry.State.EXPORTED
        ).exists()
    ):
        raise ReconciliationError("Há lançamentos já exportados; use reexportação explícita.")
    for entry in entries:
        errors = validate_journal_entry(entry)
        if errors:
            raise ReconciliationError(" ".join(errors))
    rendered = adapter.render(entries)
    digest = hashlib.sha256(rendered).hexdigest()
    export = AccountingExport(
        organization=organization,
        company=company,
        period_start=start,
        period_end=end,
        target=adapter.target,
        adapter_version=adapter.version,
        state=AccountingExport.State.READY,
        content_hash=digest,
        entry_ids=[str(entry.id) for entry in entries],
        reexport_of=reexport_of,
        reexport_reason=reexport_reason[:240],
        created_by=actor if isinstance(actor, User) else None,
    )
    export.content.save(
        f"lancamentos-{start.isoformat()}-{end.isoformat()}.csv", ContentFile(rendered), save=False
    )
    export.save()
    JournalEntry.objects.filter(id__in=[entry.id for entry in entries]).update(
        state=JournalEntry.State.EXPORTED
    )
    record_event(
        action="hub.reconciliation.export_created",
        actor=actor,
        organization=organization,
        target=export,
        request=request,
        metadata={"entries": len(entries), "content_hash": digest},
    )
    return export
