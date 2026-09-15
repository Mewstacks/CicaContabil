"""Small, scoped exports for an answer produced inside the CICA Copilot."""

from __future__ import annotations

from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.intelligence.models import Message


def _text(value: object) -> str:
    return str(value or "").replace("\x00", "").strip()


def _xlsx_text(value: object) -> str:
    """Keep spreadsheet readers from interpreting user/model content as a formula."""
    text = _text(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _evidence(message: Message) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for item in message.evidence if isinstance(message.evidence, list) else []:
        if not isinstance(item, dict):
            continue
        result.append(
            {
                "label": _text(item.get("label")),
                "reference": _text(item.get("reference")),
                "detail": _text(item.get("detail")),
            }
        )
    return result


def _metadata(message: Message) -> dict[str, str]:
    company = message.conversation.company
    if company is None:
        raise ValueError("A resposta não está vinculada a uma empresa.")
    generated_at = timezone.localtime().strftime("%d/%m/%Y %H:%M")
    return {
        "company": _text(company.name),
        "conversation": _text(message.conversation.title) or "Análise sem título",
        "generated_at": generated_at,
    }


def export_filename(*, message: Message, export_format: str) -> str:
    metadata = _metadata(message)
    base = (
        slugify(f"cica-{metadata['company']}-{metadata['conversation']}")[:72] or "cica-relatorio"
    )
    return f"{base}.{export_format}"


def build_xlsx_report(*, message: Message) -> bytes:
    metadata = _metadata(message)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Relatório CICA"
    sheet.append(["Relatório CICA"])
    sheet.append(["Empresa", _xlsx_text(metadata["company"])])
    sheet.append(["Análise", _xlsx_text(metadata["conversation"])])
    sheet.append(["Gerado em", metadata["generated_at"]])
    sheet.append([])
    sheet.append(["Resposta"])
    sheet.append([_xlsx_text(message.content)])
    sheet.append([])
    sheet.append(["Fontes", "Referência", "Detalhe"])
    for item in _evidence(message):
        sheet.append(
            [_xlsx_text(item["label"]), _xlsx_text(item["reference"]), _xlsx_text(item["detail"])]
        )
    if not _evidence(message):
        sheet.append(["Nenhuma fonte registrada nesta resposta.", "", ""])

    title_fill = PatternFill("solid", fgColor="174C3D")
    for row in (1, 6, 9):
        for cell in sheet[row]:
            cell.font = Font(bold=True, color="FFF9EE")
            cell.fill = title_fill
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 34
    sheet.column_dimensions["C"].width = 62
    sheet.freeze_panes = "A9"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _paragraph(value: object, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(_text(value)).replace("\n", "<br/>") or "—", style)


def build_pdf_report(*, message: Message) -> bytes:
    metadata = _metadata(message)
    stream = BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=f"CICA — {metadata['conversation']}",
        author="CICA · Mewstack",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("CicaTitle", parent=styles["Title"], textColor=HexColor("#174C3D"))
    heading = ParagraphStyle(
        "CicaHeading", parent=styles["Heading2"], textColor=HexColor("#174C3D")
    )
    body = ParagraphStyle("CicaBody", parent=styles["BodyText"], leading=15)
    story: list[Any] = [
        Paragraph("Relatório CICA", title),
        Spacer(1, 5 * mm),
        Table(
            [
                ["Empresa", _text(metadata["company"])],
                ["Análise", _text(metadata["conversation"])],
                ["Gerado em", metadata["generated_at"]],
            ],
            colWidths=[32 * mm, 142 * mm],
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (0, -1), HexColor("#EDF3E8")),
                    ("TEXTCOLOR", (0, 0), (0, -1), HexColor("#174C3D")),
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.4, HexColor("#C5D0C1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("PADDING", (0, 0), (-1, -1), 7),
                ]
            ),
        ),
        Spacer(1, 8 * mm),
        Paragraph("Resposta", heading),
        _paragraph(message.content, body),
        Spacer(1, 7 * mm),
        Paragraph("Fontes", heading),
    ]
    evidence = _evidence(message)
    if evidence:
        rows: list[list[Any]] = [["Fonte", "Referência", "Detalhe"]]
        rows.extend(
            [
                [
                    _paragraph(item["label"], body),
                    _paragraph(item["reference"], body),
                    _paragraph(item["detail"], body),
                ]
                for item in evidence
            ]
        )
        story.append(
            Table(
                rows,
                colWidths=[37 * mm, 47 * mm, 90 * mm],
                repeatRows=1,
                style=TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), HexColor("#174C3D")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#FFF9EE")),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.4, HexColor("#C5D0C1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("PADDING", (0, 0), (-1, -1), 6),
                    ]
                ),
            )
        )
    else:
        story.append(Paragraph("Nenhuma fonte registrada nesta resposta.", body))
    document.build(story)
    return stream.getvalue()
