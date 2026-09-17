from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from uuid import uuid4

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.hub.models import ClientCompany, OfficeProfile, ProductModule
from apps.intelligence.models import (
    AnswerFeedback,
    Conversation,
    EgressAudit,
    LearningCandidate,
    Message,
)
from apps.organizations.models import Membership, Organization
from apps.platform.models import (
    Plan,
    PlanServiceRate,
    PlatformConfiguration,
    TenantContract,
    UsageEvent,
)


class IntelligenceViewsTests(TestCase):
    databases = {"default", "knowledge"}

    def setUp(self) -> None:
        PlatformConfiguration.objects.update_or_create(
            key="default", defaults={"copilot_available_for_offices": True}
        )
        self.owner = User.objects.create_user("owner@example.test", "safe-password-123")
        self.organization = Organization.objects.create(name="Acme", slug="acme")
        OfficeProfile.objects.create(
            organization=self.organization, trial_started_at=timezone.now()
        )
        ProductModule.objects.create(organization=self.organization, code="ai", enabled=True)
        plan = Plan.objects.create(code="views-ai", name="Views AI")
        PlanServiceRate.objects.create(plan=plan, action_code="ai.answer", included_units=10)
        TenantContract.objects.create(
            organization=self.organization,
            plan=plan,
            status="trial",
            starts_on=timezone.localdate(),
            trial_ends_on=timezone.localdate() + timedelta(days=14),
            selected_modules=["ai"],
        )
        Membership.objects.create(
            organization=self.organization, user=self.owner, role=Membership.Role.OWNER
        )
        self.company = ClientCompany.objects.create(
            organization=self.organization, name="Empresa Acme", dominio_code="001"
        )
        self.client.force_login(self.owner)
        session = self.client.session
        session["hub_organization_id"] = str(self.organization.id)
        session.save()

    def assistant_message(self) -> Message:
        conversation = Conversation.objects.create(
            organization=self.organization, company=self.company
        )
        return Message.objects.create(
            organization=self.organization,
            conversation=conversation,
            role=Message.Role.ASSISTANT,
            content="Resposta rastreável.",
            evidence=[{"label": "Manual", "reference": "Procedimento", "detail": "Aprovado"}],
        )

    def test_assistant_requires_company_scope_and_keeps_conversation_in_that_company(self) -> None:
        page = self.client.get(reverse("intelligence:assistant"))
        invalid = self.client.post(reverse("intelligence:assistant"), {"question": ""})
        answered = self.client.post(
            reverse("intelligence:assistant"),
            {"question": "Quais são as pendências?", "company_id": str(self.company.id)},
        )

        self.assertEqual(page.status_code, 200)
        self.assertEqual(invalid.status_code, 200)
        self.assertContains(invalid, "Este campo é obrigatório")
        self.assertEqual(answered.status_code, 302)
        self.assertTrue(
            Conversation.objects.filter(
                organization=self.organization, company=self.company
            ).exists()
        )

    def test_same_submission_redirects_without_second_answer_or_reservation(self) -> None:
        submission_id = uuid4()
        payload = {
            "question": "Quais são as pendências?",
            "company_id": str(self.company.id),
            "request_id": str(submission_id),
        }
        first = self.client.post(reverse("intelligence:assistant"), payload)
        repeated = self.client.post(reverse("intelligence:assistant"), payload)
        self.assertEqual(first.status_code, 302)
        self.assertEqual(repeated.status_code, 302)
        self.assertEqual(first.url, repeated.url)
        self.assertEqual(Message.objects.filter(request_id=submission_id).count(), 1)
        self.assertEqual(
            Message.objects.filter(organization=self.organization, role=Message.Role.ASSISTANT)
            .exclude(in_reply_to__isnull=True)
            .count(),
            1,
        )
        self.assertEqual(UsageEvent.objects.filter(organization=self.organization).count(), 1)

    def test_company_picker_and_existing_thread_company(self) -> None:
        new_page = self.client.get(reverse("intelligence:assistant"))
        self.assertContains(new_page, 'id="assistant-company"')
        self.assertContains(new_page, "Empresa para esta análise")
        self.assertContains(new_page, self.company.name)

        conversation = Conversation.objects.create(
            organization=self.organization, company=self.company, title="Análise existente"
        )
        thread_url = reverse("intelligence:assistant") + f"?conversation={conversation.id}"
        thread_page = self.client.get(thread_url)
        self.assertContains(
            thread_page, f'name="company_id" value="{self.company.id}"'
        )
        answered = self.client.post(
            thread_url,
            {"question": "Outra dúvida da mesma empresa?", "conversation_id": str(conversation.id)},
        )
        self.assertEqual(answered.status_code, 302)
        self.assertEqual(
            Message.objects.filter(
                organization=self.organization,
                conversation=conversation,
                role=Message.Role.ASSISTANT,
            ).count(),
            1,
        )

    def test_uncertain_attempt_is_visible_beside_question_with_support_reference(self) -> None:
        PlatformConfiguration.objects.filter(key="default").update(
            support_email="suporte@example.test"
        )
        conversation = Conversation.objects.create(
            organization=self.organization, company=self.company, title="Pendência"
        )
        question = Message.objects.create(
            organization=self.organization,
            conversation=conversation,
            role=Message.Role.USER,
            content="Qual pendência?",
            request_id=uuid4(),
        )
        audit = EgressAudit.objects.create(
            organization=self.organization,
            provider="anthropic",
            purpose="technical_fallback",
            payload_hash="b" * 64,
            actor=self.owner,
            role=Membership.Role.OWNER,
            allowed=True,
            call_state=EgressAudit.CallState.UNKNOWN,
            user_message=question,
        )
        page = self.client.get(
            reverse("intelligence:assistant"), {"conversation": str(conversation.id)}
        )
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Resultado da IA incerto")
        self.assertContains(page, str(audit.id))
        self.assertContains(page, "mailto:suporte@example.test")

    def test_disabled_copilot_module_blocks_page_feedback_and_exports(self) -> None:
        message = self.assistant_message()
        feedback = AnswerFeedback.objects.create(
            organization=self.organization,
            message=message,
            submitted_by=self.owner,
            verdict=AnswerFeedback.Verdict.HELPFUL,
        )
        candidate = LearningCandidate.objects.create(
            organization=self.organization,
            feedback=feedback,
            source_summary="Fonte",
            prior_answer="Resposta anterior",
            candidate_answer="Resposta proposta",
        )
        ProductModule.objects.filter(
            organization=self.organization, code=ProductModule.Code.AI
        ).update(enabled=False)

        for url, method in (
            (reverse("intelligence:assistant"), "get"),
            (reverse("intelligence:feedback", args=[message.id]), "post"),
            (reverse("intelligence:export", args=[message.id, "pdf"]), "get"),
            (reverse("intelligence:learning"), "get"),
            (reverse("intelligence:review", args=[candidate.id, "approve"]), "post"),
        ):
            response = getattr(self.client, method)(
                url, {"verdict": "helpful"} if method == "post" else None
            )
            self.assertEqual(response.status_code, 403, url)
            self.assertContains(response, "não faz parte do acesso atual", status_code=403)

        dashboard = self.client.get(reverse("hub:dashboard"))
        self.assertNotContains(dashboard, reverse("intelligence:assistant"))
        self.assertNotContains(dashboard, "Copiloto CICA")
        candidate.refresh_from_db()
        self.assertEqual(candidate.status, LearningCandidate.Status.PENDING)

    def test_unreleased_copilot_is_hidden_and_direct_urls_are_not_available(self) -> None:
        PlatformConfiguration.objects.filter(key="default").update(
            copilot_available_for_offices=False
        )
        message = self.assistant_message()

        home = self.client.get(reverse("hub:home"))
        assistant = self.client.get(reverse("intelligence:assistant"))
        export = self.client.get(reverse("intelligence:export", args=[message.id, "pdf"]))

        self.assertNotContains(home, "Copiloto CICA")
        self.assertEqual(assistant.status_code, 404)
        self.assertEqual(export.status_code, 404)

    def test_feedback_and_learning_review_are_auditable_and_curator_gated(self) -> None:
        message = self.assistant_message()
        feedback = self.client.post(
            reverse("intelligence:feedback", args=[message.id]),
            {"verdict": AnswerFeedback.Verdict.NOT_HELPFUL, "comment": "Precisa citar a regra."},
        )
        stored_feedback = AnswerFeedback.objects.get(message=message, submitted_by=self.owner)
        candidate = LearningCandidate.objects.get(feedback=stored_feedback)

        center = self.client.get(reverse("intelligence:learning"))
        approved = self.client.post(reverse("intelligence:review", args=[candidate.id, "approve"]))
        candidate.refresh_from_db()

        self.assertRedirects(
            feedback, reverse("intelligence:assistant") + f"?conversation={message.conversation_id}"
        )
        self.assertEqual(center.status_code, 200)
        self.assertRedirects(approved, reverse("intelligence:learning"))
        self.assertEqual(candidate.status, LearningCandidate.Status.APPROVED)
        self.assertEqual(candidate.reviewed_by, self.owner)

        manager = User.objects.create_user("manager@example.test", "safe-password-123")
        Membership.objects.create(
            organization=self.organization, user=manager, role=Membership.Role.MANAGER
        )
        self.client.force_login(manager)
        session = self.client.session
        session["hub_organization_id"] = str(self.organization.id)
        session.save()
        forbidden = self.client.get(reverse("intelligence:learning"))
        self.assertEqual(forbidden.status_code, 403)

    def test_answer_exports_are_scoped_and_generate_real_pdf_and_xlsx(self) -> None:
        message = self.assistant_message()
        pdf = self.client.get(reverse("intelligence:export", args=[message.id, "pdf"]))
        xlsx = self.client.get(reverse("intelligence:export", args=[message.id, "xlsx"]))

        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertEqual(xlsx.status_code, 200)
        self.assertEqual(
            xlsx["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        sheet = load_workbook(BytesIO(xlsx.content)).active
        self.assertEqual(sheet["B2"].value, self.company.name)
        self.assertEqual(sheet["A10"].value, "Manual")

        message.content = "=SOMETHING()"
        message.save(update_fields=["content", "updated_at"])
        protected_xlsx = self.client.get(reverse("intelligence:export", args=[message.id, "xlsx"]))
        protected_sheet = load_workbook(BytesIO(protected_xlsx.content)).active
        self.assertEqual(protected_sheet["A7"].value, "'=SOMETHING()")

        other = Organization.objects.create(name="Outro", slug="outro")
        other_company = ClientCompany.objects.create(organization=other, name="Segredo")
        other_conversation = Conversation.objects.create(organization=other, company=other_company)
        other_message = Message.objects.create(
            organization=other,
            conversation=other_conversation,
            role=Message.Role.ASSISTANT,
            content="Não exportar.",
        )
        blocked = self.client.get(reverse("intelligence:export", args=[other_message.id, "pdf"]))
        self.assertEqual(blocked.status_code, 404)
