from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.db import connection
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader

from apps.accounts.models import User
from apps.audit.models import AuditEvent
from apps.hub.controlplane import company_queryset_for_membership, module_codes_for_membership
from apps.hub.models import (
    AccountingBalanceSnapshot,
    AccumulatorCatalogEntry,
    AccumulatorHistoryEntry,
    AccumulatorObservation,
    AccumulatorRule,
    Certificate,
    ClientCompany,
    CompanyAccessGrant,
    Connector,
    ControlPlaneBinding,
    DataSource,
    DreMappingSet,
    FinancialReportExport,
    ImportBatch,
    IntegrationArtifact,
    NfseExport,
    NfseSync,
    OperationalActivity,
    PayrollPeriodSnapshot,
    ProductModule,
    ReformAlert,
    ReformSourceStatus,
    ReviewCase,
)
from apps.hub.module_catalog import MODULES
from apps.hub.services import create_document_and_artifact, create_nfse_export
from apps.intelligence.models import EdgeAgent, IntelligenceConnector
from apps.organizations.models import Membership, Organization
from apps.platform.models import (
    DominioSupportTicket,
    Invitation,
    Plan,
    PlanServiceRate,
    PlatformAccess,
    SupportSession,
    TenantContract,
    TenantUsagePolicy,
)
from apps.platform.notifications import TransactionalEmailError
from apps.triage.ingest import receive_email_attachment
from apps.triage.models import (
    DestinationProfile,
    DocumentType,
    Mailbox,
    TriageEvent,
    TriageItem,
    TriageSafetyScan,
)
from apps.triage.security import ScanVerdict, scan_quarantined_item
from apps.triage.storage import PrivateTriageStorage
from apps.triage.transitions import TriageStatus
from conftest import complete_mfa


class HubWorkspaceViewTests(TestCase):
    databases = {"default", "knowledge"}

    def setUp(self) -> None:
        self.user = User.objects.create_user("owner@example.test", "safe-password-123")
        self.organization = Organization.objects.create(name="Acme", slug="acme")
        Membership.objects.create(
            organization=self.organization, user=self.user, role=Membership.Role.OWNER
        )
        self.company = ClientCompany.objects.create(
            organization=self.organization, name="Empresa Acme", dominio_code="001"
        )
        self.client.force_login(self.user)
        session = self.client.session
        session["hub_organization_id"] = str(self.organization.id)
        session.save()

    def test_the_workspace_opens_on_the_whole_portfolio_not_one_arbitrary_company(
        self,
    ) -> None:
        """An office works across its clients; nothing is pinned until somebody picks."""

        endpoints = [
            "hub:dashboard",
            "hub:nfse-center",
            "hub:companies",
            "hub:certificates",
            "hub:settings",
        ]

        for endpoint in endpoints:
            response = self.client.get(reverse(endpoint))
            self.assertEqual(response.status_code, 200, endpoint)
            self.assertNotIn("active_company", response.context, endpoint)
        reviews = self.client.get(reverse("hub:reviews"))
        self.assertRedirects(reviews, reverse("hub:nfse-center") + "?status=unclassified")

    def test_nfse_only_subscription_starts_in_nfse_and_blocks_activity_surfaces(self) -> None:
        ProductModule.objects.create(
            organization=self.organization,
            code=ProductModule.Code.NFSE,
            enabled=True,
        )
        for code in ProductModule.Code.values:
            if code != ProductModule.Code.NFSE and code != ProductModule.Code.JOURNEY:
                ProductModule.objects.create(
                    organization=self.organization,
                    code=code,
                    enabled=False,
                )
        activity = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="not-available-to-nfse-only",
            title="Atividade central",
            area="fiscal",
        )

        dashboard = self.client.get(reverse("hub:dashboard"))
        activities = self.client.get(reverse("hub:activities"))
        models = self.client.get(reverse("hub:activity-models"))
        detail = self.client.get(reverse("hub:activity-detail", args=[activity.id]))
        nfse = self.client.get(reverse("hub:nfse-center"))
        guides = self.client.get(reverse("hub:guides"))

        self.assertRedirects(dashboard, reverse("hub:nfse-center"))
        self.assertEqual(activities.status_code, 403)
        self.assertEqual(models.status_code, 403)
        self.assertEqual(detail.status_code, 404)
        self.assertEqual(guides.status_code, 403)
        self.assertContains(nfse, "NFS-e")
        self.assertContains(nfse, f'workspace-brand" href="{reverse("hub:nfse-center")}"')
        self.assertNotContains(nfse, f'href="{reverse("hub:dashboard")}"')
        self.assertNotContains(nfse, reverse("hub:activities"))
        self.assertNotContains(nfse, reverse("hub:activity-models"))

    def test_support_session_keeps_the_target_office_enabled_modules_in_navigation(self) -> None:
        """Support is tenant-scoped, but is not a collaborator with module grants."""

        ProductModule.objects.create(
            organization=self.organization,
            code=ProductModule.Code.NFSE,
            enabled=True,
        )
        support_user = User.objects.create_user("support@example.test", "safe-password-123")
        support = SupportSession.objects.create(
            organization=self.organization,
            support_user=support_user,
            justification="Verificação autorizada da configuração fiscal.",
            expires_at=timezone.now() + timedelta(minutes=20),
        )
        self.client.force_login(support_user)
        session = self.client.session
        session["hub_support_session_id"] = str(support.id)
        session.save()

        response = self.client.get(reverse("hub:certificates"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "NFS-e")
        self.assertContains(response, reverse("hub:nfse-center"))

    def test_companies_page_identifies_the_active_office(self) -> None:
        response = self.client.get(reverse("hub:companies"))

        self.assertContains(response, self.organization.name)
        self.assertContains(response, self.company.name)

    def test_setup_history_paginates_without_losing_the_selected_source(self) -> None:
        source = DataSource.objects.create(
            organization=self.organization,
            kind=DataSource.Kind.OTHER_MANUAL,
            label="Importação manual",
        )
        for index in range(21):
            ImportBatch.objects.create(
                organization=self.organization,
                data_source=source,
                kind=ImportBatch.Kind.COMPANIES,
                status=ImportBatch.Status.COMPLETED,
                original_filename=f"empresas-{index}.csv",
                content_hash=f"{index:064x}",
            )

        first_page = self.client.get(reverse("hub:setup"), {"source": source.id})
        second_page = self.client.get(
            reverse("hub:setup"), {"source": source.id, "imports_page": "2"}
        )

        self.assertEqual(first_page.status_code, 200)
        self.assertEqual(first_page.context["recent_imports_total"], 21)
        self.assertEqual(first_page.context["recent_imports_page"].number, 1)
        self.assertEqual(len(first_page.context["recent_imports"]), 20)
        self.assertContains(first_page, "21 importações")
        self.assertContains(first_page, "Página 1 de 2")
        self.assertContains(first_page, f"?source={source.id}&amp;imports_page=2")
        self.assertContains(first_page, reverse("hub:dre-mapping-editor"))
        self.assertEqual(second_page.status_code, 200)
        self.assertEqual(second_page.context["selected_source"], source)
        self.assertEqual(second_page.context["recent_imports_page"].number, 2)
        self.assertEqual(len(second_page.context["recent_imports"]), 1)
        self.assertContains(second_page, "Página 2 de 2")
        self.assertContains(second_page, f"?source={source.id}&amp;imports_page=1")

    def test_setup_is_resumable_and_links_an_administrator_to_each_pending_action(self) -> None:
        response = self.client.get(reverse("hub:setup"))

        self.assertEqual(response.status_code, 200)
        steps = response.context["setup_steps"]
        self.assertEqual(len(steps), 8)
        self.assertEqual(steps[1]["url"], f"{reverse('hub:setup')}#source-heading")
        self.assertEqual(steps[2]["url"], reverse("hub:companies"))
        self.assertEqual(steps[3]["url"], reverse("hub:activity-models"))
        self.assertEqual(steps[5]["url"], reverse("hub:team"))
        expected_mfa_url = f"{reverse('accounts:mfa-setup')}?next={reverse('hub:setup')}"
        self.assertEqual(steps[7]["url"], expected_mfa_url)
        self.assertContains(response, "Escolher fonte")
        self.assertContains(response, "Gerenciar equipe")
        self.assertContains(response, "Ativar MFA recomendado")
        self.assertContains(response, "Recomendado")

    def test_mfa_recommendation_is_available_to_an_operator(self) -> None:
        Membership.objects.filter(user=self.user, organization=self.organization).update(
            role=Membership.Role.OPERATOR
        )

        response = self.client.get(reverse("hub:setup"))

        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["can_manage_setup"])
        self.assertContains(response, "Ativar MFA recomendado")
        self.assertContains(response, reverse("accounts:mfa-setup"))

    def test_owner_confirms_office_identity_from_setup(self) -> None:
        response = self.client.post(
            reverse("hub:setup"),
            {
                "action": "save-office",
                "office-legal_name": "Acme Contábil",
                "office-cnpj": "12.345.678/0001-95",
            },
        )

        self.assertRedirects(response, reverse("hub:setup") + "#office-heading")
        profile = self.organization.hub_profile
        self.assertEqual(profile.legal_name, "Acme Contábil")
        self.assertTrue(profile.cnpj_hash)
        self.assertTrue(AuditEvent.objects.filter(action="hub.office.identity_confirmed").exists())

    def test_owner_reactivates_disabled_source_only_after_explicit_confirmation(self) -> None:
        snapshot_at = timezone.now() - timedelta(days=2)
        source = DataSource.objects.create(
            organization=self.organization,
            kind=DataSource.Kind.DOMINIO_LOCAL_AGENT,
            label="Domínio Local",
            status=DataSource.Status.DISABLED,
            capabilities=["companies"],
            source_snapshot_at=snapshot_at,
            last_error_message="Desativada por administrador.",
        )

        page = self.client.get(reverse("hub:setup"), {"source": source.id})
        self.assertContains(page, "Fonte desativada")
        self.assertContains(page, "Reativar fonte")
        self.assertContains(page, "nova sincronização")
        denied = self.client.post(
            reverse("hub:setup"),
            {"action": "reactivate-source", "source_id": source.id},
        )
        self.assertRedirects(denied, f"{reverse('hub:setup')}?source={source.id}")
        source.refresh_from_db()
        self.assertEqual(source.status, DataSource.Status.DISABLED)

        accepted = self.client.post(
            reverse("hub:setup"),
            {
                "action": "reactivate-source",
                "source_id": source.id,
                "confirm_reactivate": "yes",
            },
        )
        self.assertRedirects(accepted, f"{reverse('hub:setup')}?source={source.id}")
        source.refresh_from_db()
        self.assertEqual(source.status, DataSource.Status.NOT_CONFIGURED)
        self.assertEqual(source.capabilities, ["companies"])
        self.assertEqual(source.source_snapshot_at, snapshot_at)
        self.assertEqual(source.last_error_message, "")
        event = AuditEvent.objects.get(action="hub.data_source.reactivated")
        self.assertEqual(event.actor, self.user)
        self.assertEqual(event.metadata["previous_status"], DataSource.Status.DISABLED)
        self.assertEqual(event.metadata["current_status"], DataSource.Status.NOT_CONFIGURED)

    def test_non_administrator_cannot_reactivate_source(self) -> None:
        source = DataSource.objects.create(
            organization=self.organization,
            kind=DataSource.Kind.DOMINIO_LOCAL_AGENT,
            label="Domínio Local",
            status=DataSource.Status.DISABLED,
        )
        membership = Membership.objects.get(organization=self.organization, user=self.user)
        membership.role = Membership.Role.AUDITOR
        membership.save(update_fields=["role"])

        response = self.client.post(
            reverse("hub:setup"),
            {
                "action": "reactivate-source",
                "source_id": source.id,
                "confirm_reactivate": "yes",
            },
        )

        self.assertEqual(response.status_code, 403)
        source.refresh_from_db()
        self.assertEqual(source.status, DataSource.Status.DISABLED)

    def test_nfse_empty_states_do_not_claim_that_a_certificate_activates_capture(self) -> None:
        """A stored A1 is not a substitute for an unimplemented external NFS-e connector."""

        certificates = self.client.get(reverse("hub:certificates"))
        company = self.client.get(reverse("hub:company-detail", args=[self.company.id]))

        self.assertContains(
            certificates, "a primeira coleta da empresa entra na fila automaticamente"
        )
        self.assertNotContains(certificates, "ativar a consulta de NFS-e")
        self.assertContains(company, "conexão NFS-e desta empresa ser homologada")
        self.assertNotContains(company, "certificado não há consulta de NFS-e nem DTE")

    def test_company_detail_centralizes_activities_and_authorized_triage_items(self) -> None:
        activity = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="folha-fechamento",
            title="Conferir fechamento da folha",
            area="payroll",
            work_status=OperationalActivity.WorkStatus.BLOCKED,
            processing_status=OperationalActivity.ProcessingStatus.OPEN,
            obligation_status=OperationalActivity.ObligationStatus.NOT_PREPARED,
            freshness=OperationalActivity.Freshness.CURRENT,
        )
        item = TriageItem.objects.create(
            organization=self.organization,
            company=self.company,
            original_name="folha-setembro.pdf",
            status=TriageStatus.AWAITING_REVIEW,
        )

        without_triage = self.client.get(reverse("hub:company-detail", args=[self.company.id]))

        self.assertEqual(without_triage.status_code, 200)
        self.assertContains(without_triage, activity.title)
        self.assertContains(without_triage, "Trabalho e fechamento")
        self.assertNotContains(without_triage, "Documentos em triagem")

        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.TRIAGE, enabled=True
        )
        response = self.client.get(reverse("hub:company-detail", args=[self.company.id]))

        self.assertContains(response, "Documentos em triagem")
        self.assertContains(response, item.original_name)
        self.assertContains(response, reverse("hub:activity-detail", args=[activity.id]))
        self.assertContains(response, reverse("hub:triage-item", args=[item.id]))

    def test_company_detail_orders_activities_by_the_effective_due_date(self) -> None:
        today = timezone.localdate()
        legal_overdue = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="legal-overdue",
            title="Prazo legal vencido",
            area="fiscal",
            legal_due_on=today - timedelta(days=1),
        )
        internal_future = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="internal-future",
            title="Prazo interno futuro",
            area="fiscal",
            legal_due_on=today - timedelta(days=5),
            internal_due_on=today + timedelta(days=3),
        )
        without_due_date = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="without-due-date",
            title="Sem prazo",
            area="fiscal",
        )

        response = self.client.get(reverse("hub:company-detail", args=[self.company.id]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            [item.id for item in response.context["activities"]],
            [legal_overdue.id, internal_future.id, without_due_date.id],
        )

    def test_company_detail_names_people_and_notes_instead_of_technical_ids(self) -> None:
        """D-277: the accountant reads names, fiscal numbers and actionable states."""

        operator = User.objects.create_user(
            "ana.interna@example.test", "safe-password-123", full_name="Ana Martins"
        )
        Membership.objects.create(
            organization=self.organization, user=operator, role=Membership.Role.OPERATOR
        )
        OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="fiscal-fechamento",
            title="Conferir apuração",
            area="fiscal",
            assigned_to=operator,
            freshness=OperationalActivity.Freshness.CURRENT,
        )
        document, _artifact, _review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='fiscal-number' />",
            normalized_data={"number": "4321", "counterparty_name": "Fornecedor Sul"},
            source_nsu="000000000077",
        )

        response = self.client.get(reverse("hub:company-detail", args=[self.company.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ana Martins")
        self.assertNotContains(response, "ana.interna@example.test")
        self.assertContains(response, "4321")
        self.assertContains(response, "Fornecedor Sul")
        self.assertNotContains(response, "000000000077")
        self.assertNotContains(response, document.document_hash[:12])
        self.assertNotContains(response, "Não verificado · Não aplicável")

    def test_team_workload_shows_roles_in_portuguese(self) -> None:
        operator = User.objects.create_user(
            "bruno@example.test", "safe-password-123", full_name="Bruno Costa"
        )
        Membership.objects.create(
            organization=self.organization, user=operator, role=Membership.Role.OPERATOR
        )

        response = self.client.get(reverse("hub:dashboard"), {"view": "management"})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Operador")
        self.assertNotContains(response, ">Operator<")
        self.assertNotContains(response, ">Owner<")

    def test_company_detail_shows_aggregate_payroll_source_without_claiming_official_query(
        self,
    ) -> None:
        PayrollPeriodSnapshot.objects.create(
            organization=self.organization,
            company=self.company,
            competence=timezone.localdate().replace(day=1),
            source_kind=PayrollPeriodSnapshot.SourceKind.DOCUMENT,
            source_reference="resumo-folha-setembro",
            workforce_count=12,
            gross_pay_cents=123_450,
            net_pay_cents=101_200,
        )

        response = self.client.get(reverse("hub:company-detail", args=[self.company.id]))

        self.assertContains(response, "Conferência da folha")
        self.assertContains(response, "resumo-folha-setembro")
        self.assertContains(response, "R$ 1.234,50")
        self.assertContains(response, "Descontos")
        self.assertContains(response, "Encargos")
        self.assertContains(
            response, "<dt>Descontos</dt><dd><span>Não informado</span>", html=True
        )
        self.assertContains(
            response, "<dt>Encargos</dt><dd><span>Não informado</span>", html=True
        )
        self.assertNotContains(response, "Nenhum valor será calculado")

    def test_company_detail_compares_payroll_sources_for_one_competence(self) -> None:
        competence = timezone.localdate().replace(day=1)
        OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="FOLHA-CONFERIR",
            title="Tratar divergencia da folha",
            area="payroll",
            competence=competence,
        )
        baseline = PayrollPeriodSnapshot.objects.create(
            organization=self.organization,
            company=self.company,
            competence=competence,
            source_kind=PayrollPeriodSnapshot.SourceKind.ERP,
            source_reference="folha-erp",
            workforce_count=10,
            gross_pay_cents=100_000,
        )
        comparison = PayrollPeriodSnapshot.objects.create(
            organization=self.organization,
            company=self.company,
            competence=competence,
            source_kind=PayrollPeriodSnapshot.SourceKind.DOCUMENT,
            source_reference="recibo-folha",
            workforce_count=12,
            gross_pay_cents=101_500,
        )
        PayrollPeriodSnapshot.objects.create(
            organization=self.organization,
            company=self.company,
            competence=(competence - timedelta(days=1)).replace(day=1),
            source_kind=PayrollPeriodSnapshot.SourceKind.MANUAL,
            source_reference="outra-competencia",
            workforce_count=9,
        )

        overview = self.client.get(reverse("hub:company-detail", args=[self.company.id]))
        self.assertContains(overview, "Conferir competência")
        self.assertNotContains(overview, 'name="left_snapshot"')

        response = self.client.get(
            reverse("hub:company-detail", args=[self.company.id]),
            {
                "left_snapshot": baseline.id,
                "right_snapshot": comparison.id,
                "money_tolerance": "0.00",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Conferir diferenças entre fontes")
        self.assertContains(response, "Quadro de pessoas")
        self.assertContains(response, "Total bruto")
        self.assertContains(response, "R$ 15,00")
        self.assertContains(response, "a mais")
        self.assertContains(response, "não calcula a folha")
        self.assertNotContains(response, "outra-competencia")
        comparison_sources = response.context["payroll_comparison_form"].fields[
            "left_snapshot"
        ].queryset
        self.assertEqual(comparison_sources.count(), 2)
        self.assertContains(
            response,
            f"company={self.company.id}&amp;area=payroll&amp;competence={competence:%Y-%m}",
            html=False,
        )

        activity_queue = self.client.get(
            reverse("hub:activities"),
            {"company": self.company.id, "area": "payroll", "competence": f"{competence:%Y-%m}"},
        )
        self.assertContains(activity_queue, "Tratar divergencia da folha")
        invalid_competence = self.client.get(reverse("hub:activities"), {"competence": "2030-13"})
        self.assertContains(
            invalid_competence, "Informe uma competência válida no formato mês/ano."
        )

    def test_certificate_workspace_prioritizes_expiry_and_missing_companies(self) -> None:
        valid_company = ClientCompany.objects.create(
            organization=self.organization,
            name="Empresa com certificado longo",
            dominio_code="090",
        )
        expiring_company = ClientCompany.objects.create(
            organization=self.organization,
            name="Empresa vencendo logo",
            dominio_code="091",
        )
        Certificate.objects.create(
            organization=self.organization,
            company=valid_company,
            label="A1 matriz",
            pfx_blob="encrypted-valid",
            password="encrypted-password",
            fingerprint_sha256="a" * 64,
            valid_until=timezone.now() + timedelta(days=90),
        )
        Certificate.objects.create(
            organization=self.organization,
            company=expiring_company,
            label="A1 filial",
            pfx_blob="encrypted-expiring",
            password="encrypted-password",
            fingerprint_sha256="b" * 64,
            valid_until=timezone.now() + timedelta(days=5),
        )

        attention = self.client.get(reverse("hub:certificates"))

        self.assertContains(attention, "Empresa vencendo logo")
        self.assertContains(attention, "A1 matriz")
        self.assertEqual(attention.context["certificate_status"], "all")
        self.assertContains(attention, "Vence em breve")
        self.assertContains(attention, "Empresas sem A1 válido")
        self.assertContains(attention, self.company.name)

        searched = self.client.get(reverse("hub:certificates"), {"q": "090", "status": "all"})
        self.assertContains(searched, "Empresa com certificado longo")
        self.assertEqual(
            [certificate.company for certificate in searched.context["certificates"]],
            [valid_company],
        )
        self.assertContains(searched, "1 resultado")

    def test_certificate_workspace_paginates_missing_companies_without_hiding_them(self) -> None:
        for index in range(20):
            ClientCompany.objects.create(
                organization=self.organization,
                name=f"Empresa sem A1 {index:03d}",
                dominio_code=f"A{index:03d}",
            )

        first_page = self.client.get(reverse("hub:certificates"), {"q": "matriz", "status": "all"})
        second_page = self.client.get(
            reverse("hub:certificates"),
            {"q": "matriz", "status": "all", "missing_page": "2"},
        )

        self.assertContains(first_page, "21 empresas")
        self.assertContains(first_page, "Página 1 de 2")
        self.assertContains(first_page, "Empresa sem A1 000")
        self.assertNotIn(
            "Empresa sem A1 019",
            [company.name for company in first_page.context["missing_certificate_companies"]],
        )
        self.assertContains(second_page, "Página 2 de 2")
        self.assertContains(second_page, "Empresa sem A1 019")
        self.assertEqual(
            [company.name for company in second_page.context["missing_certificate_companies"]],
            ["Empresa sem A1 019"],
        )
        self.assertContains(
            second_page,
            "?q=matriz&amp;status=all&amp;missing_page=1#empresas-sem-certificado",
            html=False,
        )

    def test_demo_certificate_action_is_session_only_and_never_stores_a_pfx(self) -> None:
        self.organization.is_demo = True
        self.organization.save(update_fields=["is_demo"])
        session = self.client.session
        session["demo_visit_id"] = "certificate-demo-visitor"
        session.save()

        response = self.client.post(
            reverse("hub:certificates"), {"demo_company_id": str(self.company.id)}, follow=True
        )

        self.assertContains(response, "Certificado fictício registrado")
        self.assertContains(response, "Nenhum arquivo ou senha foi recebido")
        self.assertEqual(response.context["certificate_stats"]["missing_companies"], 0)
        self.assertFalse(Certificate.objects.filter(organization=self.organization).exists())
        self.assertTrue(
            self.client.session["demo_progress"]["certificates"][str(self.company.id)]["simulated"]
        )

    def test_reform_radar_filters_official_alerts_without_an_extra_page(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.REFORM, enabled=True
        )
        ReformAlert.objects.create(
            source=ReformAlert.Source.RFB,
            external_key="rfb-1",
            title="Cronograma IBS",
            source_url="https://www.gov.br/receitafederal/noticia",
            relevance=ReformAlert.Relevance.REFORM,
            content_hash="hash-1",
        )
        ReformAlert.objects.create(
            source=ReformAlert.Source.FAZENDA,
            external_key="fazenda-1",
            title="Crédito tributário da CBS",
            source_url="https://www.gov.br/fazenda/noticia",
            relevance=ReformAlert.Relevance.FISCAL,
            content_hash="hash-2",
        )

        response = self.client.get(reverse("hub:reform"), {"fonte": "rfb"})

        self.assertContains(response, "Cronograma IBS")
        self.assertNotContains(response, "Crédito tributário da CBS")
        self.assertContains(response, "Aplicar filtros")
        self.assertContains(response, "Situação das fontes oficiais")
        self.assertContains(response, "Aguardando a primeira coleta")

        unfiltered = self.client.get(reverse("hub:reform"))
        self.assertContains(unfiltered, "Cronograma IBS")
        self.assertContains(unfiltered, "Crédito tributário da CBS")
        self.assertNotContains(unfiltered, 'class="operational-table"')
        self.assertContains(unfiltered, "Analisar impacto", count=2)
        self.assertContains(unfiltered, "Abrir fonte oficial", count=2)
        self.assertContains(unfiltered, "?return_to=", count=2)

        searched = self.client.get(reverse("hub:reform"), {"q": "crédito", "relevancia": "fiscal"})
        self.assertContains(searched, "Crédito tributário da CBS")
        self.assertNotContains(searched, "Cronograma IBS")
        self.assertContains(searched, "1 publicação")

        membership = Membership.objects.get(organization=self.organization, user=self.user)
        membership.role = Membership.Role.AUDITOR
        membership.save(update_fields=["role"])
        readonly = self.client.get(reverse("hub:reform"))
        self.assertContains(readonly, "Ver análises", count=2)
        self.assertNotContains(readonly, "Analisar impacto")

    def test_reform_radar_paginates_the_full_filtered_history(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.REFORM, enabled=True
        )
        for index in range(101):
            ReformAlert.objects.create(
                source=ReformAlert.Source.RFB,
                external_key=f"radar-pagination-{index}",
                title=f"Alerta Radar paginado {index:03d}",
                source_url=f"https://www.gov.br/receitafederal/noticia-{index}",
                relevance=ReformAlert.Relevance.REFORM,
                content_hash=f"hash-radar-{index}",
            )

        first_page = self.client.get(
            reverse("hub:reform"), {"fonte": "rfb", "relevancia": "reform"}
        )
        last_page = self.client.get(
            reverse("hub:reform"),
            {"fonte": "rfb", "relevancia": "reform", "page": "3"},
        )

        self.assertEqual(first_page.status_code, 200)
        self.assertEqual(first_page.context["alert_total"], 101)
        self.assertEqual(first_page.context["alert_page"].number, 1)
        self.assertEqual(len(first_page.context["alerts"]), 50)
        self.assertEqual(last_page.context["alert_page"].number, 3)
        self.assertContains(last_page, "Alerta Radar paginado 000")
        self.assertContains(last_page, "101 publicações")
        self.assertContains(last_page, "Página 3 de 3")
        self.assertContains(
            last_page,
            "?fonte=rfb&amp;relevancia=reform&amp;page=2",
            html=False,
        )

    def test_reform_radar_shows_a_source_collection_failure_without_raw_error(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.REFORM, enabled=True
        )
        ReformSourceStatus.objects.create(
            source=ReformAlert.Source.RFB,
            last_error="upstream timeout with internal trace 9d1f",
        )

        response = self.client.get(reverse("hub:reform"))

        self.assertContains(response, "Falha na coleta")
        self.assertContains(response, "a próxima coleta tentará novamente")
        self.assertNotContains(response, "internal trace 9d1f")

    def test_account_menu_persists_an_explicit_theme_choice(self) -> None:
        response = self.client.post(
            reverse("hub:set-theme"), {"theme": "dark", "next": reverse("hub:companies")}
        )

        self.assertRedirects(response, reverse("hub:companies"))
        self.assertEqual(response.cookies["hub_theme"].value, "dark")
        page = self.client.get(reverse("hub:companies"))
        self.assertContains(page, 'data-theme="dark"')
        self.assertContains(page, 'aria-pressed="true"')

    def test_account_menu_rejects_an_unknown_theme_choice(self) -> None:
        response = self.client.post(reverse("hub:set-theme"), {"theme": "rainbow"})

        self.assertEqual(response.status_code, 400)

    @patch("apps.hub.views.send_invitation_email")
    def test_owner_invites_a_collaborator_with_company_and_module_scope(self, send_email) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.TRIAGE, enabled=True
        )
        response = self.client.post(
            reverse("hub:team"),
            {
                "full_name": "Ana Fiscal",
                "email": "ana@example.test",
                "role": Membership.Role.OPERATOR,
                "modules": [ProductModule.Code.NFSE, ProductModule.Code.TRIAGE],
                "companies": [str(self.company.id)],
            },
        )

        self.assertRedirects(response, reverse("hub:team"))
        invitation = Invitation.objects.get(email="ana@example.test")
        self.assertEqual(invitation.company_ids, [str(self.company.id)])
        self.assertEqual(
            set(invitation.modules), {ProductModule.Code.NFSE, ProductModule.Code.TRIAGE}
        )
        send_email.assert_called_once()

    def test_invitation_validation_marks_a_real_focus_target(self) -> None:
        response = self.client.post(reverse("hub:team"), {"role": Membership.Role.MANAGER})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "data-focus-error", count=4)
        self.assertNotContains(response, 'data-focus-error"=""')

    @patch(
        "apps.hub.views.send_invitation_email",
        side_effect=TransactionalEmailError("offline"),
    )
    def test_invitation_is_not_created_when_delivery_fails(self, _send_email) -> None:
        response = self.client.post(
            reverse("hub:team"),
            {
                "full_name": "Ana Fiscal",
                "email": "ana@example.test",
                "role": Membership.Role.OPERATOR,
                "modules": [ProductModule.Code.NFSE],
                "companies": [str(self.company.id)],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Não foi possível enviar o convite")
        self.assertFalse(Invitation.objects.filter(email="ana@example.test").exists())

    @patch("apps.hub.views.send_invitation_email")
    def test_owner_can_resend_a_pending_collaborator_invitation(self, send_email) -> None:
        token, digest = Invitation.issue_token()
        invitation = Invitation.objects.create(
            organization=self.organization,
            email="ana@example.test",
            role=Membership.Role.OPERATOR,
            company_ids=[str(self.company.id)],
            modules=[ProductModule.Code.NFSE],
            token_digest=digest,
            expires_at=timezone.now() - timedelta(hours=1),
        )

        response = self.client.post(
            reverse("hub:collaborator-invitation-resend", args=[invitation.id])
        )

        self.assertRedirects(response, reverse("hub:team"))
        invitation.refresh_from_db()
        self.assertNotEqual(invitation.token_digest, digest)
        self.assertTrue(invitation.usable())
        send_email.assert_called_once()
        self.assertNotIn(token, str(send_email.call_args))

    def test_owner_can_revoke_a_pending_collaborator_invitation(self) -> None:
        _token, digest = Invitation.issue_token()
        invitation = Invitation.objects.create(
            organization=self.organization,
            email="ana@example.test",
            role=Membership.Role.OPERATOR,
            company_ids=[str(self.company.id)],
            modules=[ProductModule.Code.NFSE],
            token_digest=digest,
            expires_at=timezone.now() + timedelta(days=1),
        )

        response = self.client.post(
            reverse("hub:collaborator-invitation-revoke", args=[invitation.id])
        )

        self.assertRedirects(response, reverse("hub:team"))
        invitation.refresh_from_db()
        self.assertEqual(invitation.status, Invitation.Status.REVOKED)
        self.assertFalse(invitation.usable())

    def test_collaborator_acceptance_limits_companies_and_modules(self) -> None:
        other_company = ClientCompany.objects.create(
            organization=self.organization, name="Outra empresa", dominio_code="002"
        )
        token, digest = Invitation.issue_token()
        Invitation.objects.create(
            organization=self.organization,
            email="operador@example.test",
            full_name="Operador",
            role=Membership.Role.OPERATOR,
            company_ids=[str(self.company.id)],
            modules=[ProductModule.Code.NFSE],
            token_digest=digest,
            expires_at=timezone.now() + timedelta(days=1),
        )

        response = self.client.post(
            reverse("hub:activate", args=[token]),
            {"password": "a-safe-password-123", "password_confirm": "a-safe-password-123"},
        )

        self.assertEqual(response.status_code, 302)
        invited = User.objects.get(email="operador@example.test")
        membership = Membership.objects.get(organization=self.organization, user=invited)
        self.assertEqual(list(company_queryset_for_membership(membership)), [self.company])
        self.assertEqual(module_codes_for_membership(membership), {ProductModule.Code.NFSE})
        grant = CompanyAccessGrant.objects.get(membership=membership, company=self.company)
        self.assertEqual(grant.modules, [ProductModule.Code.NFSE])
        self.assertFalse(
            CompanyAccessGrant.objects.filter(membership=membership, company=other_company).exists()
        )

    def test_collaborator_cannot_bypass_module_scope_by_opening_a_url(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.GUIDES, enabled=True
        )
        collaborator = User.objects.create_user(
            email="scoped@example.test", password="a-safe-password-123"
        )
        membership = Membership.objects.create(
            organization=self.organization, user=collaborator, role=Membership.Role.OPERATOR
        )
        CompanyAccessGrant.objects.create(
            organization=self.organization,
            membership=membership,
            company=self.company,
            modules=[ProductModule.Code.NFSE],
            capabilities=["*"],
        )
        self.client.force_login(collaborator)
        session = self.client.session
        session["hub_organization_id"] = str(self.organization.id)
        session.save()

        self.assertEqual(self.client.get(reverse("hub:nfse-center")).status_code, 200)
        self.assertEqual(self.client.get(reverse("hub:guides")).status_code, 403)
        self.assertEqual(self.client.get(reverse("hub:team")).status_code, 403)

    def test_owner_can_request_an_update_only_when_the_dominio_agent_is_online(self) -> None:
        connector = IntelligenceConnector.objects.create(
            organization=self.organization,
            mode=IntelligenceConnector.Mode.EDGE_AGENT,
            status="healthy",
        )
        EdgeAgent.objects.create(
            organization=self.organization,
            label="Servidor Domínio",
            fingerprint="agent-online",
            shared_secret="shared-secret",
            last_seen_at=timezone.now(),
        )

        page = self.client.get(reverse("hub:settings"))
        self.assertContains(page, "Atualizar agora")
        self.assertContains(page, "Aguardando primeira atualização")

        response = self.client.post(reverse("hub:settings"), {"action": "request-dominio-sync"})

        connector.refresh_from_db()
        self.assertRedirects(response, reverse("hub:settings"))
        self.assertIsNotNone(connector.sync_requested_at)

    def test_direct_dominio_dsn_is_encrypted_at_rest(self) -> None:
        connector = IntelligenceConnector.objects.create(
            organization=self.organization,
            mode=IntelligenceConnector.Mode.DIRECT_ODBC,
            odbc_dsn="DSN=Contabil;UID=readonly;PWD=not-a-real-password",
        )
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT odbc_dsn FROM intelligence_intelligenceconnector WHERE id = %s",
                [connector._meta.pk.get_db_prep_value(connector.id, connection=connection)],
            )
            stored_value = cursor.fetchone()[0]
        self.assertNotIn("readonly", stored_value)
        self.assertNotIn("not-a-real-password", stored_value)
        self.assertTrue(stored_value.startswith("enc:v1:"))
        connector.refresh_from_db()
        self.assertEqual(
            connector.odbc_dsn,
            "DSN=Contabil;UID=readonly;PWD=not-a-real-password",
        )

    @patch("apps.hub.views.ReadOnlyDominoOdbc")
    def test_owner_can_update_a_direct_dominio_connection_now(self, mocked_odbc) -> None:
        connector = IntelligenceConnector.objects.create(
            organization=self.organization,
            mode=IntelligenceConnector.Mode.DIRECT_ODBC,
            status="healthy",
            odbc_dsn="Contabil",
        )
        mocked_odbc.return_value.execute.side_effect = lambda query: (
            [
                {
                    "codigo": "001",
                    "nome": "Empresa Acme",
                    "cnpj_masked": "12.345.678/0001-90",
                }
            ]
            if query == "companies"
            else []
        )

        self.assertContains(self.client.get(reverse("hub:settings")), "Atualizar agora")
        response = self.client.post(
            reverse("hub:settings"), {"action": "request-dominio-sync"}, follow=True
        )

        self.company.refresh_from_db()
        connector.refresh_from_db()
        mocked_odbc.assert_called_once_with("Contabil")
        self.assertEqual(
            [call.args for call in mocked_odbc.return_value.execute.call_args_list],
            [("companies",), ("bank_entries",)],
        )
        self.assertEqual(self.company.cnpj_masked, "12.345.678/0001-90")
        self.assertContains(response, "Domínio atualizado: 1 empresa(s) e 0 item(ns) bancários.")
        self.assertIsNone(connector.sync_requested_at)

    def test_direct_dominio_connection_is_ready_for_dominio_modules(self) -> None:
        IntelligenceConnector.objects.create(
            organization=self.organization,
            mode=IntelligenceConnector.Mode.DIRECT_ODBC,
            status="healthy",
            odbc_dsn="Contabil",
        )
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.RECONCILIATION, enabled=True
        )

        response = self.client.get(reverse("hub:reconciliation"))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["connector_ready"])

    def test_settings_exposes_a_bank_snapshot_limit_warning(self) -> None:
        IntelligenceConnector.objects.create(
            organization=self.organization,
            mode=IntelligenceConnector.Mode.DIRECT_ODBC,
            status="healthy",
            last_sync_at=timezone.now(),
            odbc_dsn="Contabil",
        )
        AuditEvent.objects.create(
            organization=self.organization,
            action="intelligence.dominio.bank_entries_synced",
            metadata={"may_be_truncated": True},
        )

        response = self.client.get(reverse("hub:settings"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Cobertura parcial de itens bancarios.")
        self.assertTrue(response.context["dominio_bank_entries_may_be_truncated"])

    def test_reconciliation_imports_only_through_the_upload_dialog(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.RECONCILIATION, enabled=True
        )

        response = self.client.post(reverse("hub:reconciliation"), {"company": self.company.pk})

        self.assertEqual(response.status_code, 405)
        page = self.client.get(reverse("hub:reconciliation"))
        self.assertContains(page, 'id="reconciliation-upload-dialog"')
        self.assertNotContains(page, 'id="import-ofx"')

    @patch("apps.hub.views.ReadOnlyDominoOdbc")
    def test_a_failed_direct_sync_offers_resolution_and_a_developer_ticket(
        self, mocked_odbc
    ) -> None:
        connector = IntelligenceConnector.objects.create(
            organization=self.organization,
            mode=IntelligenceConnector.Mode.DIRECT_ODBC,
            status="healthy",
            odbc_dsn="Contabil",
        )
        mocked_odbc.return_value.execute.side_effect = RuntimeError("ODBC indisponível")

        self.client.post(reverse("hub:settings"), {"action": "request-dominio-sync"})
        connector.refresh_from_db()
        page = self.client.get(reverse("hub:settings"))
        response = self.client.post(
            reverse("hub:settings"), {"action": "dominio-support-ticket"}, follow=True
        )

        self.assertEqual(connector.status, "error")
        self.assertEqual(connector.last_error_code, "odbc_sync")
        self.assertContains(page, "Atenção")
        self.assertContains(page, "Resolver")
        self.assertContains(response, "Chamado aberto para desenvolvimento.")
        self.assertTrue(
            DominioSupportTicket.objects.filter(
                organization=self.organization, connector=connector
            ).exists()
        )

    def test_enabled_product_modules_render_for_the_office(self) -> None:
        for code in (
            ProductModule.Code.GUIDES,
            ProductModule.Code.INTEGRA,
            ProductModule.Code.RECONCILIATION,
            ProductModule.Code.REFORM,
        ):
            ProductModule.objects.create(organization=self.organization, code=code, enabled=True)

        for endpoint, heading in (
            ("hub:guides", "Guias e DCTFWeb"),
            ("hub:reconciliation", "Conciliação OFX x Domínio"),
            ("hub:reform", "Radar da Reforma Tributária"),
        ):
            response = self.client.get(reverse(endpoint))
            self.assertEqual(response.status_code, 200, endpoint)
            self.assertContains(response, heading)
            self.assertContains(response, self.organization.name)

        integra = self.client.get(reverse("hub:integra"))
        self.assertEqual(integra.status_code, 200)
        self.assertContains(integra, "Caixa DTE")
        self.assertContains(integra, "Parcelamentos")
        self.assertContains(integra, "DCTFWeb")
        self.assertContains(integra, reverse("hub:dte-center"))

        nav = self.client.get(reverse("hub:dashboard"))
        self.assertContains(nav, reverse("hub:guides"))
        self.assertContains(nav, reverse("hub:integra"))

    def test_triage_module_requires_an_email_mailbox_when_enabled(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.TRIAGE, enabled=True
        )

        response = self.client.get(reverse("hub:triage"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Triagem de Arquivos")
        self.assertContains(response, "Caixa de e-mail não configurada")
        self.assertContains(response, reverse("hub:triage-connections"))
        self.assertNotContains(response, "Pronto para receber dados")

        connections = self.client.get(reverse("hub:triage-connections"))
        self.assertEqual(connections.status_code, 200)
        self.assertContains(connections, "Nenhuma caixa de e-mail conectada")

        nav = self.client.get(reverse("hub:dashboard"))
        self.assertContains(nav, reverse("hub:triage"))

    def test_triage_item_routes_respect_the_collaborators_company_scope(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.TRIAGE, enabled=True
        )
        restricted_company = ClientCompany.objects.create(
            organization=self.organization, name="Empresa restrita", dominio_code="002"
        )
        restricted_item = TriageItem.objects.create(
            organization=self.organization,
            company=restricted_company,
            original_name="documento.pdf",
        )
        collaborator = User.objects.create_user(
            email="triage-scoped@example.test", password="a-safe-password-123"
        )
        membership = Membership.objects.create(
            organization=self.organization, user=collaborator, role=Membership.Role.OPERATOR
        )
        CompanyAccessGrant.objects.create(
            organization=self.organization,
            membership=membership,
            company=self.company,
            modules=[ProductModule.Code.TRIAGE],
            capabilities=["read"],
        )
        self.client.force_login(collaborator)
        session = self.client.session
        session["hub_organization_id"] = str(self.organization.id)
        session.save()

        for endpoint in ("hub:triage-item", "hub:triage-download"):
            self.assertEqual(
                self.client.get(reverse(endpoint, args=[restricted_item.id])).status_code,
                404,
            )

    def test_triage_queue_shows_unassigned_mail_only_to_admin_and_filters_by_stage(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.TRIAGE, enabled=True
        )
        unknown = TriageItem.objects.create(
            organization=self.organization,
            original_name="pendente-de-identificacao.pdf",
            status=TriageItem.Status.QUARANTINED,
        )
        assigned = TriageItem.objects.create(
            organization=self.organization,
            company=self.company,
            original_name="documento-atribuido.xml",
            status=TriageItem.Status.AWAITING_REVIEW,
        )
        response = self.client.get(reverse("hub:triage"))
        self.assertContains(response, unknown.original_name)
        self.assertContains(response, assigned.original_name)
        self.assertContains(response, "1 em quarentena")
        self.assertEqual(
            self.client.get(reverse("hub:triage-item", args=[unknown.id])).status_code, 200
        )
        filtered = self.client.get(reverse("hub:triage"), {"status": TriageItem.Status.QUARANTINED})
        self.assertContains(filtered, unknown.original_name)
        self.assertNotContains(filtered, assigned.original_name)

        collaborator = User.objects.create_user(
            email="triage-queue-scope@example.test", password="a-safe-password-123"
        )
        membership = Membership.objects.create(
            organization=self.organization, user=collaborator, role=Membership.Role.OPERATOR
        )
        CompanyAccessGrant.objects.create(
            organization=self.organization,
            membership=membership,
            company=self.company,
            modules=[ProductModule.Code.TRIAGE],
            capabilities=["read"],
        )
        self.client.force_login(collaborator)
        session = self.client.session
        session["hub_organization_id"] = str(self.organization.id)
        session.save()
        operator_page = self.client.get(reverse("hub:triage"))
        self.assertContains(operator_page, assigned.original_name)
        self.assertNotContains(operator_page, unknown.original_name)
        self.assertEqual(
            self.client.get(reverse("hub:triage-item", args=[unknown.id])).status_code, 404
        )

    def test_triage_prototype_cannot_review_or_download_unscanned_files(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.TRIAGE, enabled=True
        )
        item = TriageItem.objects.create(
            organization=self.organization,
            company=self.company,
            original_name="anexo-nao-verificado.pdf",
        )

        detail = self.client.get(reverse("hub:triage-item", args=[item.id]))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "verificação de segurança")
        self.assertNotContains(detail, "Baixar arquivo")
        refused = self.client.post(reverse("hub:triage-item", args=[item.id]), {}, follow=True)
        self.assertContains(refused, "Escolha uma decisão válida")
        self.assertEqual(
            self.client.get(reverse("hub:triage-download", args=[item.id])).status_code,
            404,
        )

    def test_triage_item_history_pages_all_events_without_a_silent_cutoff(self) -> None:
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.TRIAGE, enabled=True
        )
        item = TriageItem.objects.create(
            organization=self.organization,
            company=self.company,
            original_name="historico-completo.pdf",
        )
        for index in range(21):
            TriageEvent.objects.create(
                organization=self.organization,
                triage_item=item,
                actor=self.user,
                from_status=TriageStatus.RECEIVED,
                to_status=TriageStatus.QUARANTINED,
                note=f"Evento de auditoria {index:02d}",
            )

        first_page = self.client.get(
            reverse("hub:triage-item", args=[item.id]), {"return_to": "triagem"}
        )
        second_page = self.client.get(
            reverse("hub:triage-item", args=[item.id]), {"history_page": 2}
        )

        self.assertEqual(first_page.status_code, 200)
        self.assertEqual(second_page.status_code, 200)
        self.assertEqual(first_page.context["triage_history_total"], 21)
        self.assertEqual(first_page.context["triage_history_page"].number, 1)
        self.assertEqual(second_page.context["triage_history_page"].number, 2)
        self.assertEqual(len(first_page.context["triage_history_events"]), 20)
        self.assertEqual(len(second_page.context["triage_history_events"]), 1)
        self.assertContains(first_page, "Evento de auditoria 00")
        self.assertNotContains(first_page, "Evento de auditoria 20")
        self.assertContains(second_page, "Evento de auditoria 20")
        self.assertContains(first_page, "Página 1 de 2 · 21 eventos")
        self.assertContains(second_page, "Página 2 de 2 · 21 eventos")
        self.assertContains(first_page, "return_to=triagem&amp;history_page=2")
        self.assertContains(second_page, "history_page=1")

    def test_internal_library_download_uses_verified_copy_and_refuses_tampering(self) -> None:
        self.enterContext(override_settings(MEDIA_ROOT=self.enterContext(TemporaryDirectory())))
        ProductModule.objects.create(
            organization=self.organization, code=ProductModule.Code.TRIAGE, enabled=True
        )
        mailbox = Mailbox.objects.create(
            organization=self.organization,
            provider=Mailbox.Provider.IMAP,
            address="arquivo@acme.test",
            active=True,
            status=Mailbox.Status.ACTIVE,
        )
        payload = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\n%%EOF\n"
        item = receive_email_attachment(
            mailbox=mailbox,
            message_id="download-test",
            part_id="1",
            filename="documento.pdf",
            payload=payload,
        ).item

        class CleanScanner:
            def scan(self, stream: object) -> ScanVerdict:
                return ScanVerdict(TriageSafetyScan.Verdict.CLEAN, "test-scanner")

        scan_quarantined_item(item=item, scanner=CleanScanner())
        item.refresh_from_db()
        doc_type = DocumentType.objects.create(
            organization=self.organization,
            code="documento",
            label="Documento",
            name_template="{codigo}_DOCUMENTO_{periodo}",
            period_kind=DocumentType.PeriodKind.COMPETENCIA,
        )
        DestinationProfile.objects.create(
            organization=self.organization, mode=DestinationProfile.Mode.INTERNAL
        )
        item.company = self.company
        item.document_type = doc_type
        item.final_name = "001_DOCUMENTO_082026.pdf"
        item.transition_to(TriageStatus.EXTRACTING)
        item.transition_to(TriageStatus.AWAITING_REVIEW)
        item.save()
        detail_url = reverse("hub:triage-item", args=[item.id])
        self.assertContains(self.client.get(detail_url), "Aprovar para arquivamento")
        approved = self.client.post(detail_url, {"decision": "archive"}, follow=True)
        self.assertContains(approved, "Arquivar na biblioteca interna")
        item.refresh_from_db()
        archived = self.client.post(detail_url, {"decision": "finish_archive"}, follow=True)
        self.assertContains(archived, "Baixar arquivo da biblioteca")
        item.refresh_from_db()
        detail = self.client.get(reverse("hub:triage-item", args=[item.id]))
        self.assertContains(detail, "Baixar arquivo da biblioteca")
        self.assertContains(detail, "Aguardando extração")
        self.assertContains(detail, "Cópia interna íntegra confirmada")
        download = self.client.get(reverse("hub:triage-download", args=[item.id]))
        self.assertEqual(download.status_code, 200)
        self.assertEqual(b"".join(download.streaming_content), payload)
        self.assertEqual(download["Cache-Control"], "no-store, private")
        self.assertIn("attachment", download["Content-Disposition"])
        self.assertTrue(download.closed)
        with PrivateTriageStorage().open(item.destination_path, "wb") as copy:
            copy.write(b"tampered")
        self.assertEqual(
            self.client.get(reverse("hub:triage-download", args=[item.id])).status_code,
            404,
        )

    def test_triage_module_is_unavailable_until_the_office_enables_it(self) -> None:
        response = self.client.get(reverse("hub:triage"))

        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "não está habilitado", status_code=403)
        self.assertNotContains(self.client.get(reverse("hub:dashboard")), reverse("hub:triage"))

    def test_disabled_product_module_is_not_available(self) -> None:
        ProductModule.objects.create(
            organization=self.organization,
            code=ProductModule.Code.RECONCILIATION,
            enabled=False,
        )
        response = self.client.get(reverse("hub:reconciliation"))
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "não está habilitado", status_code=403)
        self.assertNotContains(
            self.client.get(reverse("hub:dashboard")), reverse("hub:reconciliation")
        )

    def test_retired_journey_routes_and_navigation_are_unavailable(self) -> None:
        ProductModule.objects.create(
            organization=self.organization,
            code=ProductModule.Code.JOURNEY,
            enabled=True,
        )
        for path in (
            "/app/jornada/",
            "/app/jornada/etapas/criar/",
            "/app/jornada/pedidos/criar/",
            "/app/jornada/etapas/concluir/",
            "/app/jornada/pedidos/transicionar/",
        ):
            self.assertEqual(self.client.get(path).status_code, 404, path)
            self.assertEqual(self.client.post(path, {}).status_code, 404, path)
        self.assertNotIn(ProductModule.Code.JOURNEY, MODULES)
        self.assertNotContains(self.client.get(reverse("hub:dashboard")), "/app/jornada/")

    def test_unavailable_external_connector_rejects_credentials(self) -> None:
        response = self.client.post(
            reverse("hub:settings"),
            {
                "kind": Connector.Kind.ONVIO,
                "label": "Onvio QA",
                "endpoint": "",
                "database_alias": "DOMINIO_QA",
                "secret": "local-secret",
            },
        )
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "Não informe credenciais.", status_code=403)
        self.assertFalse(
            Connector.objects.filter(
                organization=self.organization, kind=Connector.Kind.ONVIO
            ).exists()
        )

    def test_legacy_per_call_usage_rule_is_rejected_in_favor_of_tokens(
        self,
    ) -> None:
        plan = Plan.objects.create(code="usage", name="Uso", monthly_price_cents=9_900)
        PlanServiceRate.objects.create(
            plan=plan,
            action_code="caixapostal.mensagens",
            included_units=50,
            overage_unit_price_cents=125,
        )
        contract = TenantContract.objects.create(
            organization=self.organization, plan=plan, status=TenantContract.Status.ACTIVE
        )
        rate = contract.service_rates.get(action_code="caixapostal.mensagens")
        prefix = f"usage_{rate.id}"
        complete_mfa(self.client)

        response = self.client.post(
            reverse("hub:settings"),
            {
                "action": "usage-policy",
                "rate_id": rate.id,
                f"{prefix}-overage_mode": TenantUsagePolicy.OverageMode.ALLOW,
                f"{prefix}-warning_percent": 80,
                f"{prefix}-monthly_overage_cap_brl": "35.50",
            },
        )

        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "cobrança por chamada foi descontinuada", status_code=403)
        self.assertFalse(TenantUsagePolicy.objects.filter(organization=self.organization).exists())
        self.assertFalse(
            Connector.objects.filter(organization=self.organization, kind=Connector.Kind.INTEGRA)
        )

        page = self.client.get(reverse("hub:settings"))
        self.assertNotContains(page, "chamadas incluídas")
        self.assertNotContains(page, "por chamada")
        self.assertContains(page, "proposta com valor do token")

    def test_generic_connector_posts_are_not_accepted(self) -> None:
        response = self.client.post(
            reverse("hub:settings"),
            {"kind": Connector.Kind.DOMINIO_AGENT, "label": "Domínio", "secret": "x"},
        )

        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "ainda não está disponível", status_code=403)
        self.assertFalse(
            Connector.objects.filter(
                organization=self.organization, kind=Connector.Kind.DOMINIO_AGENT
            ).exists()
        )

    def test_settings_survives_a_legacy_connector_with_an_unavailable_key(self) -> None:
        connector = Connector.objects.create(
            organization=self.organization,
            kind=Connector.Kind.INTEGRA,
            enabled=True,
            status="configured",
            encrypted_configuration="initial configuration",
        )
        with connection.cursor() as cursor:
            cursor.execute(
                "UPDATE hub_connector SET encrypted_configuration = %s WHERE id = %s",
                [
                    "enc:v1:test-v1:AAAAAAAAAAAAAAAA:AAAAAAAAAAAAAAAAAAAAAA==",
                    connector._meta.pk.get_db_prep_value(connector.id, connection=connection),
                ],
            )

        response = self.client.get(reverse("hub:settings"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Integra")

    def test_nfse_center_shows_the_whole_portfolio(self) -> None:
        other_company = ClientCompany.objects.create(
            organization=self.organization, name="Outra empresa", dominio_code="002"
        )
        create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='owned' />",
            normalized_data={"number": "owned"},
            source_nsu="owned",
        )
        create_document_and_artifact(
            company=other_company,
            original_xml="<nfse id='other' />",
            normalized_data={"number": "other"},
            source_nsu="other",
        )

        response = self.client.get(
            reverse("hub:nfse-center"), {"date_filter": "competence", "competence_month": ""}
        )

        self.assertContains(response, "<h1>NFS-e</h1>", html=True)
        self.assertContains(response, "owned")
        self.assertContains(response, "other")

        scoped = self.client.get(reverse("hub:company-detail", args=[self.company.id]))

        self.assertContains(scoped, "owned")
        self.assertNotContains(scoped, "other")

    def test_nfse_center_indicators_follow_the_list_filters(self) -> None:
        other_company = ClientCompany.objects.create(
            organization=self.organization, name="Outra empresa", dominio_code="003"
        )
        for company in (self.company, other_company):
            for month in ("08", "09"):
                for index in range(2):
                    create_document_and_artifact(
                        company=company,
                        original_xml=f"<nfse id='{company.pk}-{month}-{index}' />",
                        normalized_data={
                            "number": f"{month}{index}",
                            "issued_at": f"2026-{month}-10T12:00:00-03:00",
                        },
                        source_nsu=f"{company.pk}-{month}-{index}",
                    )

        response = self.client.get(
            reverse("hub:nfse-center"),
            {
                "company": self.company.pk,
                "status": "all",
                "date_filter": "competence",
                "competence_month": "09",
                "competence_year": "2026",
            },
        )

        stats = response.context["nfse_stats"]
        self.assertEqual(stats["received"], 2)
        self.assertEqual(stats["pending"] + stats["classified"], 2)
        self.assertEqual(stats["provided"] + stats["taken"] + stats["unknown"], 2)
        for url in response.context["nfse_stat_urls"].values():
            self.assertIn("competence_month=09", url)
            self.assertIn("competence_year=2026", url)
            self.assertIn(f"company={self.company.pk}", url)

    def test_nfse_center_states_classification_instead_of_a_score(self) -> None:
        create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='classificacao' />",
            normalized_data={},
            source_nsu="classificacao",
        )

        response = self.client.get(
            reverse("hub:nfse-center"), {"date_filter": "competence", "competence_month": ""}
        )

        self.assertContains(response, '<th scope="col">Acumulador</th>', html=True)
        self.assertNotContains(response, "Confiança")
        self.assertNotContains(response, "Transitória · 0%")
        body = response.content.decode()
        self.assertIn("Classificar", body)
        self.assertIn("Ver dados da nota", body)

    def test_nfse_center_shows_note_number_and_not_transport_identifiers(self) -> None:
        document, _artifact, _review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='visible-number' />",
            normalized_data={"number": "NF-1042"},
            source_nsu="NSU-TECHNICAL-99",
        )

        response = self.client.get(
            reverse("hub:nfse-center"), {"date_filter": "competence", "competence_month": ""}
        )

        self.assertContains(response, "NF-1042")
        self.assertNotContains(response, "NSU-TECHNICAL-99")
        self.assertNotContains(response, document.document_hash[:12])

    def test_nfse_retention_reports_follow_filters_and_export_explicit_taxes(self) -> None:
        common = {
            "competence": "2026-09-01",
            "service_code": "1401",
            "service_description": "Assessoria contábil mensal",
            "amount": "1000.00",
            "issued_at": "2026-09-10T12:00:00-03:00",
        }
        create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='provided-report' />",
            normalized_data={
                **common,
                "number": "SAIDA-1",
                "direction": "provided",
                "counterparty_name": "Cliente sem retenção",
            },
            source_nsu="provided-report",
        )
        create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='taken-report' />",
            normalized_data={
                **common,
                "number": "ENTRADA-1",
                "direction": "taken",
                "counterparty_name": "=FORNECEDOR",
                "retentions": {
                    "iss": "20.00",
                    "pis": "6.50",
                    "cofins": "30.00",
                    "csll": "10.00",
                    "irrf": "15.00",
                    "inss": "110.00",
                },
                "retained_total": "9999.99",
            },
            source_nsu="taken-report",
        )
        filters = {
            "report_status": "all",
            "report_direction": "taken",
            "report_date_filter": "competence",
            "report_competence": "2026-09",
        }

        xlsx = self.client.post(
            reverse("hub:nfse-center") + "?status=all",
            {"action": "retention_report_xlsx", **filters},
        )

        self.assertEqual(xlsx.status_code, 200)
        self.assertEqual(
            xlsx["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(xlsx.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Resumo", "Notas"])
        notes = workbook["Notas"]
        self.assertEqual(notes.max_row, 2)
        self.assertEqual(notes["C2"].value, "ENTRADA-1")
        self.assertEqual(notes["D2"].value, "Entrada · serviço tomado")
        self.assertEqual(notes["G2"].value, "'=FORNECEDOR")
        self.assertEqual(notes["K2"].value, 20)
        self.assertEqual(notes["Q2"].value, "=SUM(K2:P2)")
        self.assertEqual(workbook["Resumo"]["B11"].value, "=SUM(Notas!Q2:Q2)")

        pdf = self.client.post(
            reverse("hub:nfse-center") + "?status=all",
            {"action": "retention_report_pdf", **filters},
        )
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        pdf_text = "\n".join(
            page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages
        )
        self.assertIn("ENTRADA-1", pdf_text)
        self.assertNotIn("SAIDA-1", pdf_text)
        self.assertIn("R$ 191,50", pdf_text)
        self.assertEqual(
            AuditEvent.objects.filter(
                organization=self.organization,
                action__in={
                    "hub.nfse.retention_report_pdf_downloaded",
                    "hub.nfse.retention_report_xlsx_downloaded",
                },
            ).count(),
            2,
        )

    def test_classified_nfse_accumulator_can_be_corrected_without_rewriting_history(self) -> None:
        document, _artifact, review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='correct-accumulator' />",
            normalized_data={"number": "NF-204", "service_code": "1401"},
            source_nsu="204",
        )
        assert review is not None
        previous = IntegrationArtifact.objects.create(
            organization=self.organization,
            document=document,
            accumulator_code="AC-ANTIGO",
            applied_rule="Decisão anterior",
            confidence=100,
        )
        AccumulatorRule.objects.create(
            organization=self.organization,
            company=self.company,
            name="Novo acumulador",
            accumulator_code="AC-NOVO",
            priority=10,
            match={"service_code": "1401"},
        )

        response = self.client.post(
            reverse("hub:nfse-update-accumulator", args=[document.id]),
            {"accumulator_code": "AC-NOVO", "return_to": reverse("hub:nfse-center")},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
            HTTP_ACCEPT="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["accumulator"], "AC-NOVO")
        artifacts = list(
            IntegrationArtifact.objects.filter(document=document).order_by("created_at")
        )
        self.assertEqual(len(artifacts), 2)
        self.assertEqual(artifacts[0].id, previous.id)
        self.assertEqual(artifacts[-1].accumulator_code, "AC-NOVO")
        review.refresh_from_db()
        self.assertEqual(review.status, ReviewCase.Status.RESOLVED)
        self.assertEqual(review.resolved_accumulator, "AC-NOVO")
        self.assertTrue(
            AccumulatorHistoryEntry.objects.filter(
                organization=self.organization,
                company=self.company,
                accumulator_code="AC-NOVO",
            ).exists()
        )

    def test_nfse_bulk_export_selects_all_classified_results_from_the_current_filter(self) -> None:
        other_company = ClientCompany.objects.create(
            organization=self.organization, name="Empresa incluída", dominio_code="002"
        )
        excluded_company = ClientCompany.objects.create(
            organization=self.organization, name="Empresa fora do filtro", dominio_code="003"
        )
        included, _, _ = create_document_and_artifact(
            company=other_company,
            original_xml=(
                "<NFSe><infNFSe><numero>included</numero>"
                "<valores><vLiq>1</vLiq></valores></infNFSe></NFSe>"
            ),
            normalized_data={"number": "NF-INCLUIDA"},
            source_nsu="included",
        )
        excluded, _, _ = create_document_and_artifact(
            company=excluded_company,
            original_xml=(
                "<NFSe><infNFSe><numero>excluded</numero>"
                "<valores><vLiq>1</vLiq></valores></infNFSe></NFSe>"
            ),
            normalized_data={"number": "NF-EXCLUIDA"},
            source_nsu="excluded",
        )
        for document, accumulator in ((included, "AC-2"), (excluded, "AC-3")):
            IntegrationArtifact.objects.create(
                organization=self.organization,
                document=document,
                accumulator_code=accumulator,
                applied_rule="Teste",
                confidence=100,
            )

        response = self.client.post(
            reverse("hub:nfse-center"),
            {
                "action": "export_dominio_xml",
                "all_filtered_classified": "1",
                "export_q": "Empresa incluída",
                "export_date_filter": "competence",
                "export_competence": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/zip")
        export = NfseExport.objects.get(organization=self.organization)
        self.assertEqual(export.document_count, 1)
        self.assertEqual(export.snapshot["documents"][0]["document_id"], str(included.id))

    def test_nfse_company_link_and_bulk_export_use_exact_company_not_its_name(self) -> None:
        other = ClientCompany.objects.create(
            organization=self.organization, name=self.company.name, dominio_code="OTHER"
        )
        for company in (self.company, other):
            document, _, _ = create_document_and_artifact(
                company=company,
                original_xml=(
                    f"<NFSe company='{company.pk}'><infNFSe>"
                    "<valores><vLiq>1</vLiq></valores></infNFSe></NFSe>"
                ),
                normalized_data={"number": str(company.pk)},
                source_nsu=str(company.pk),
            )
            IntegrationArtifact.objects.create(
                organization=self.organization,
                document=document,
                accumulator_code="A1",
                applied_rule="Test",
                confidence=100,
            )
        url = reverse("hub:nfse-center") + f"?company={self.company.pk}&status=all"
        page = self.client.get(url)
        self.assertEqual(page.context["document_filtered_total"], 1)
        self.assertEqual(page.context["nfse_selected_company"], self.company)
        response = self.client.post(
            url, {"action": "export_dominio_xml", "all_filtered_classified": "1"}
        )
        self.assertEqual(response.status_code, 200)
        export = NfseExport.objects.get(organization=self.organization)
        self.assertEqual(export.document_count, 1)
        self.assertEqual(
            self.client.get(reverse("hub:nfse-center"), {"company": "invalid"}).status_code,
            404,
        )
        foreign = Organization.objects.create(name="Foreign", slug="foreign-nfse-filter")
        forbidden = ClientCompany.objects.create(organization=foreign, name=self.company.name)
        self.assertEqual(
            self.client.get(reverse("hub:nfse-center"), {"company": forbidden.pk}).status_code,
            404,
        )

    def test_paused_company_nfse_history_and_download_preserve_access_boundaries(self) -> None:
        document, _, _ = create_document_and_artifact(
            company=self.company,
            original_xml=(
                "<NFSe id='paused-history'><infNFSe>"
                "<valores><vLiq>1</vLiq></valores></infNFSe></NFSe>"
            ),
            normalized_data={"number": "PAUSED-1"},
            source_nsu="paused-history",
        )
        IntegrationArtifact.objects.create(
            organization=self.organization,
            document=document,
            accumulator_code="A1",
            applied_rule="Test",
            confidence=100,
        )
        export = create_nfse_export(
            organization=self.organization, documents=[document], actor=self.user
        )
        self.company.active = False
        self.company.save(update_fields=["active"])
        url = reverse("hub:nfse-center") + f"?company={self.company.pk}&status=all"
        page = self.client.get(url)
        self.assertContains(page, "HISTÓRICO SOMENTE PARA CONSULTA")
        self.assertContains(page, f"{self.company.name} está pausada")
        self.assertContains(page, "Nenhuma nota nova será coletada")
        self.assertContains(page, "Ver situação no cadastro")
        self.assertContains(page, "Ver pendências")
        self.assertContains(page, "Inclui 1 nota do filtro atual")
        self.assertNotContains(page, "1 para classificar")
        self.assertFalse(page.context["can_classify_nfse"])
        self.assertNotContains(page, "data-nfse-accumulator-input")
        self.assertEqual(page.context["document_filtered_total"], 1)
        portfolio = self.client.get(reverse("hub:nfse-center"), {"status": "all"})
        self.assertEqual(portfolio.context["document_filtered_total"], 0)
        download_url = reverse("hub:nfse-export-download", args=[export.id])
        response = self.client.post(download_url)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(b"".join(response.streaming_content).startswith(b"PK"))
        generated = self.client.post(
            url, {"action": "export_dominio_xml", "all_filtered_classified": "1"}
        )
        self.assertEqual(generated.status_code, 200)
        self.assertTrue(b"".join(generated.streaming_content).startswith(b"PK"))
        self.client.post(url, {"action": "activate", "companies": [str(self.company.pk)]})
        self.company.refresh_from_db()
        self.assertFalse(self.company.active)
        self.assertFalse(NfseSync.objects.filter(company=self.company).exists())

        membership = Membership.objects.get(organization=self.organization, user=self.user)
        membership.role = Membership.Role.OPERATOR
        membership.save(update_fields=["role"])
        CompanyAccessGrant.objects.create(
            organization=self.organization,
            membership=membership,
            company=self.company,
            modules=[ProductModule.Code.NFSE],
            is_active=True,
        )
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.post(download_url).status_code, 404)

    def test_nfse_bulk_export_rejects_invalid_filters_without_broadening_selection(self) -> None:
        document, _, _ = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='strict-export' />",
            normalized_data={"number": "STRICT-1"},
            source_nsu="strict-export",
        )
        IntegrationArtifact.objects.create(
            organization=self.organization,
            document=document,
            accumulator_code="AC-1",
            applied_rule="Teste",
            confidence=100,
        )
        cases = [
            {"export_date_filter": "issued", "export_issued_from": "31/02/2026"},
            {"export_date_filter": "issued", "export_issued_to": "invalid"},
            {
                "export_date_filter": "issued",
                "export_issued_from": "30/09/2026",
                "export_issued_to": "01/09/2026",
            },
            {"export_date_filter": "competence", "export_competence": "2026-13"},
            {"export_date_filter": "unknown"},
        ]
        for filters in cases:
            with self.subTest(filters=filters):
                response = self.client.post(
                    reverse("hub:nfse-center") + "?status=classified",
                    {"action": "export_dominio_xml", "all_filtered_classified": "1", **filters},
                )
                self.assertRedirects(
                    response,
                    reverse("hub:nfse-center") + "?status=classified",
                    fetch_redirect_response=False,
                )
                self.assertFalse(NfseExport.objects.filter(organization=self.organization).exists())
        for invalid_id in ("not-a-uuid", "00000000-0000-0000-0000-000000000001"):
            with self.subTest(invalid_id=invalid_id):
                response = self.client.post(
                    reverse("hub:nfse-center"),
                    {"action": "export_dominio_xml", "documents": [str(document.id), invalid_id]},
                )
                self.assertEqual(response.status_code, 302)
                self.assertFalse(NfseExport.objects.filter(organization=self.organization).exists())

    @patch("apps.hub.views.timezone.localdate", return_value=date(2026, 9, 29))
    def test_nfse_center_defaults_to_previous_competence(self, _localdate: object) -> None:
        _august, _artifact, _review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='august' />",
            normalized_data={"number": "august", "issued_at": "2026-08-15T12:00:00-03:00"},
            source_nsu="august",
        )
        _september, _artifact, _review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='september' />",
            normalized_data={
                "number": "september",
                "issued_at": "2026-09-01T12:00:00-03:00",
            },
            source_nsu="september",
        )
        response = self.client.get(reverse("hub:nfse-center"))

        self.assertContains(response, "august")
        self.assertNotContains(response, "september")
        self.assertContains(response, '<option value="08" selected>Agosto</option>', html=True)
        self.assertContains(response, 'value="2026"')

    def test_nfse_center_is_an_operational_bulk_sync_workspace(self) -> None:
        response = self.client.get(reverse("hub:nfse-center"), {"view": "collection"})

        self.assertContains(response, "Situação da coleta")
        self.assertContains(response, "Andamento por empresa")
        self.assertContains(response, "Configurar empresas")
        self.assertContains(response, reverse("hub:nfse-queue-status"))
        self.assertContains(response, "Selecionar todas as empresas exibidas")
        self.assertContains(response, "Ativar coleta")
        self.assertNotContains(response, "Coleta externa ainda desligada neste ambiente")
        self.assertContains(response, self.company.name)

    def test_nfse_live_queue_is_scoped_and_expired_certificates_are_skipped(self) -> None:
        valid_company = ClientCompany.objects.create(
            organization=self.organization, name="Empresa válida", dominio_code="010"
        )
        Certificate.objects.create(
            organization=self.organization,
            company=valid_company,
            label="A1 válido",
            pfx_blob="encrypted-valid",
            password="encrypted-password",
            fingerprint_sha256="c" * 64,
            valid_until=timezone.now() + timedelta(days=90),
        )
        NfseSync.objects.create(
            organization=self.organization,
            company=valid_company,
            enabled=True,
            status=NfseSync.Status.RUNNING,
            last_run_at=timezone.now(),
        )

        response = self.client.get(reverse("hub:nfse-queue-status"))

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        states = {item["company"]: item["state"] for item in payload["items"]}
        self.assertEqual(states["Empresa válida"], "running")
        self.assertEqual(states[self.company.name], "blocked")
        self.assertEqual(payload["counts"]["running"], 1)
        self.assertEqual(payload["counts"]["skipped"], 1)
        self.assertEqual(payload["counts"]["requires_action"], 1)
        self.assertFalse(payload["simulated"])
        self.assertEqual(response["Cache-Control"], "private, no-store")

    def test_nfse_live_queue_retries_only_failed_companies_without_page_reload(self) -> None:
        certificate = Certificate.objects.create(
            organization=self.organization,
            company=self.company,
            label="A1 válido",
            pfx_blob="encrypted-valid",
            password="encrypted-password",
            fingerprint_sha256="d" * 64,
            valid_until=timezone.now() + timedelta(days=90),
        )
        sync = NfseSync.objects.create(
            organization=self.organization,
            company=self.company,
            certificate=certificate,
            enabled=True,
            status=NfseSync.Status.ERROR,
            last_error_code="adn_error",
            last_error_message="Falha anterior",
        )

        with (
            self.captureOnCommitCallbacks(execute=True),
            patch("apps.hub.tasks.dispatch_active_nfse_syncs.delay") as dispatch,
        ):
            response = self.client.post(
                reverse("hub:nfse-queue-retry"), {"company_id": str(self.company.id)}
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"changed": 1})
        sync.refresh_from_db()
        self.assertEqual(sync.status, NfseSync.Status.RETRY)
        self.assertEqual(sync.last_error_message, "")
        self.assertIsNotNone(sync.next_run_at)
        dispatch.assert_called_once_with()

    def test_nfse_activation_without_a_valid_company_certificate_is_blocked(self) -> None:
        response = self.client.post(
            reverse("hub:nfse-center"),
            {"action": "activate", "companies": [str(self.company.id)]},
            follow=True,
        )

        self.assertRedirects(response, reverse("hub:nfse-center"))
        self.assertFalse(
            NfseSync.objects.filter(organization=self.organization, company=self.company).exists()
        )
        self.assertContains(response, "ignorada por não ter e-CNPJ válido")
        self.assertContains(response, "continuam normalmente")

    def test_nfse_bulk_pause_stops_scheduling_without_removing_history(self) -> None:
        sync = NfseSync.objects.create(
            organization=self.organization,
            company=self.company,
            enabled=True,
            status=NfseSync.Status.IDLE,
            checkpoint_nsu="981",
            next_run_at=timezone.now(),
        )

        response = self.client.post(
            reverse("hub:nfse-center"),
            {"action": "pause", "companies": [str(self.company.id)]},
        )

        self.assertRedirects(response, reverse("hub:nfse-center"))
        sync.refresh_from_db()
        self.assertFalse(sync.enabled)
        self.assertEqual(sync.status, NfseSync.Status.PAUSED)
        self.assertEqual(sync.checkpoint_nsu, "981")

    def test_nfse_center_shows_the_whole_portfolio_when_no_company_is_chosen(self) -> None:
        other_company = ClientCompany.objects.create(
            organization=self.organization, name="Outra empresa", dominio_code="002"
        )
        create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='owned' />",
            normalized_data={"number": "owned"},
            source_nsu="owned",
        )
        create_document_and_artifact(
            company=other_company,
            original_xml="<nfse id='other' />",
            normalized_data={"number": "other"},
            source_nsu="other",
        )

        response = self.client.get(
            reverse("hub:nfse-center"), {"date_filter": "competence", "competence_month": ""}
        )

        self.assertContains(response, "owned")
        self.assertContains(response, "other")

    def test_nfse_clear_filters_really_removes_the_default_competence(self) -> None:
        response = self.client.get(reverse("hub:nfse-center"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            f'href="{reverse("hub:nfse-center")}?status=all" data-live-filter-clear',
        )

    def test_nfse_center_paginates_the_portfolio_without_hiding_documents(self) -> None:
        now = timezone.now()
        for index in range(101):
            create_document_and_artifact(
                company=self.company,
                original_xml=f"<nfse id='page-{index}' />",
                normalized_data={
                    "number": f"PAGE-{index:03d}",
                    "service_code": "unmatched",
                    "issued_at": (now - timedelta(days=index)).isoformat(),
                },
                source_nsu=f"PAGE-{index:03d}",
            )

        first_page = self.client.get(reverse("hub:nfse-center"), {"q": "PAGE"})
        second_page = self.client.get(reverse("hub:nfse-center"), {"q": "PAGE", "page": "2"})

        self.assertContains(first_page, "101 resultados")
        self.assertContains(first_page, "Página 1 de 2")
        self.assertContains(first_page, "PAGE-000")
        self.assertNotContains(first_page, "PAGE-100")
        self.assertContains(second_page, "Página 2 de 2")
        self.assertContains(second_page, "PAGE-100")
        self.assertContains(second_page, "?q=PAGE&amp;page=1", html=False)

    def test_nfse_center_searches_the_portfolio_and_links_the_exact_review(self) -> None:
        other_company = ClientCompany.objects.create(
            organization=self.organization, name="Empresa pesquisável", dominio_code="0888"
        )
        _document, _artifact, review = create_document_and_artifact(
            company=other_company,
            original_xml="<nfse id='search-review' />",
            normalized_data={"number": "NOTA-PESQUISA", "service_code": "sem-regra"},
            source_nsu="NSU-PESQUISA",
        )
        assert review is not None

        response = self.client.get(reverse("hub:nfse-center"), {"q": "0888", "status": "review"})

        self.assertContains(response, other_company.name)
        self.assertContains(response, "NOTA-PESQUISA")
        self.assertNotContains(response, "NSU-PESQUISA")
        self.assertContains(response, reverse("hub:resolve-review", args=[review.id]))
        self.assertNotContains(response, self.company.name)

    def test_logout_uses_post_and_ends_the_workspace_session(self) -> None:
        response = self.client.post(reverse("hub:logout"))

        self.assertRedirects(response, reverse("hub:home"))
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_platform_user_gets_a_return_link_in_the_account_menu(self) -> None:
        PlatformAccess.objects.create(user=self.user, role=PlatformAccess.Role.DEVELOPER)
        complete_mfa(self.client)

        response = self.client.get(reverse("hub:dashboard"))

        self.assertContains(response, reverse("platform:dashboard"))
        self.assertContains(response, "Voltar à plataforma")

    def test_company_and_office_switches_are_scoped_to_active_membership(self) -> None:
        second_company = ClientCompany.objects.create(
            organization=self.organization, name="Outra empresa", dominio_code="002"
        )
        other_organization = Organization.objects.create(name="Outra", slug="outra")
        Membership.objects.create(
            organization=other_organization, user=self.user, role=Membership.Role.MANAGER
        )

        office_response = self.client.post(
            reverse("hub:switch-office"), {"organization_id": str(other_organization.id)}
        )

        self.assertRedirects(office_response, reverse("hub:dashboard"))
        self.assertEqual(self.client.session["hub_organization_id"], str(other_organization.id))
        portfolio = self.client.get(reverse("hub:companies")).context["companies"]
        self.assertNotIn(second_company, portfolio)

    def test_company_creation_is_available_only_to_unmanaged_installations(self) -> None:
        response = self.client.post(
            reverse("hub:companies"), {"name": "Criada no Hub", "dominio_code": "003"}
        )
        self.assertRedirects(response, reverse("hub:companies"))
        self.assertTrue(
            ClientCompany.objects.filter(
                organization=self.organization, dominio_code="003"
            ).exists()
        )

        ControlPlaneBinding.objects.create(
            organization=self.organization,
            remote_installation_id="5b3dd36b-4da4-4b7a-9f12-8e933c2ed0e4",
            controller_url="https://crmew.example.test",
            device_private_key="local-key",
            controller_public_key="controller-key",
            cache_expires_at=timezone.now() + timedelta(minutes=10),
        )
        blocked = self.client.post(
            reverse("hub:companies"), {"name": "Nunca criada", "dominio_code": "004"}
        )
        self.assertEqual(blocked.status_code, 403)
        self.assertFalse(
            ClientCompany.objects.filter(
                organization=self.organization, dominio_code="004"
            ).exists()
        )

    def test_review_resolution_requires_an_accumulator_and_records_a_human_decision(self) -> None:
        document, _artifact, review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='pending' />",
            normalized_data={"service_code": "unmatched"},
        )
        self.assertIsNotNone(document)
        assert review is not None
        AccumulatorRule.objects.create(
            organization=self.organization,
            company=self.company,
            name="Acumulador revisado",
            accumulator_code="AC-200",
            priority=1,
        )
        url = reverse("hub:resolve-review", args=[review.id])

        missing = self.client.post(url, {})
        resolved = self.client.post(url, {"accumulator_code": "AC-200"})

        detail_url = reverse("hub:review-detail", args=[review.id])
        self.assertRedirects(missing, detail_url)
        self.assertRedirects(resolved, detail_url)
        review.refresh_from_db()
        self.assertEqual(review.status, ReviewCase.Status.RESOLVED)
        self.assertEqual(review.resolved_accumulator, "AC-200")
        self.assertEqual(review.resolved_by, self.user)
        observation = AccumulatorObservation.objects.get(
            organization=self.organization, company=self.company, accumulator_code="AC-200"
        )
        self.assertEqual(observation.frequency, 1)
        self.assertTrue(
            IntegrationArtifact.objects.filter(
                organization=self.organization,
                document=document,
                accumulator_code="AC-200",
            ).exists()
        )

    def test_review_resolution_rejects_unknown_or_other_company_accumulator(self) -> None:
        _document, _artifact, review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='catalog' />",
            normalized_data={"service_code": "unmatched"},
        )
        assert review is not None
        other_company = ClientCompany.objects.create(
            organization=self.organization, name="Outra empresa", dominio_code="099"
        )
        AccumulatorRule.objects.create(
            organization=self.organization,
            company=other_company,
            name="Somente outra empresa",
            accumulator_code="AC-OUTRA",
            priority=1,
        )
        AccumulatorObservation.objects.create(
            organization=self.organization,
            company=self.company,
            accumulator_code="AC-HISTORICO",
            last_used_at=timezone.now(),
        )
        backup_source = DataSource.objects.create(
            organization=self.organization,
            kind=DataSource.Kind.DOMINIO_WEB_BACKUP,
            label="Domínio Web",
        )
        backup_batch = ImportBatch.objects.create(
            organization=self.organization,
            data_source=backup_source,
            kind=ImportBatch.Kind.DOMINIO_BACKUP,
            status=ImportBatch.Status.COMPLETED,
            original_filename="fotografia.dom",
            content_hash="a" * 64,
            source_snapshot_at=timezone.now(),
        )
        AccumulatorCatalogEntry.objects.create(
            organization=self.organization,
            company=self.company,
            data_source=backup_source,
            source_batch=backup_batch,
            accumulator_code="AC-BACKUP",
            name="Servico importado",
            source_snapshot_at=timezone.now(),
        )
        url = reverse("hub:resolve-review", args=[review.id])

        rejected = self.client.post(url, {"accumulator_code": "AC-OUTRA"})
        review.refresh_from_db()
        rejected_detail = self.client.get(reverse("hub:review-detail", args=[review.id]))
        accepted = self.client.post(url, {"accumulator_code": "AC-BACKUP"})

        self.assertRedirects(rejected, reverse("hub:review-detail", args=[review.id]))
        self.assertEqual(review.status, ReviewCase.Status.OPEN)
        self.assertContains(
            rejected_detail,
            "Escolha um acumulador cadastrado para esta empresa",
        )
        self.assertRedirects(accepted, reverse("hub:review-detail", args=[review.id]))
        review.refresh_from_db()
        self.assertEqual(review.resolved_accumulator, "AC-BACKUP")

    def test_nfse_catalog_is_scoped_to_the_company_portfolio(self) -> None:
        source = DataSource.objects.create(
            organization=self.organization,
            kind=DataSource.Kind.DOMINIO_WEB_BACKUP,
            label="Domínio Web",
        )
        batch = ImportBatch.objects.create(
            organization=self.organization,
            data_source=source,
            kind=ImportBatch.Kind.DOMINIO_BACKUP,
            status=ImportBatch.Status.COMPLETED,
            original_filename="fotografia.dom",
            content_hash="b" * 64,
            source_snapshot_at=timezone.now(),
        )
        AccumulatorCatalogEntry.objects.create(
            organization=self.organization,
            company=self.company,
            data_source=source,
            source_batch=batch,
            accumulator_code="AC-001",
            name="Serviços prestados",
            source_snapshot_at=timezone.now(),
        )
        other_company = ClientCompany.objects.create(
            organization=self.organization, name="Fora da carteira", dominio_code="999"
        )
        AccumulatorCatalogEntry.objects.create(
            organization=self.organization,
            company=other_company,
            data_source=source,
            source_batch=batch,
            accumulator_code="AC-999",
            source_snapshot_at=timezone.now(),
        )
        operator = User.objects.create_user("catalog@example.test", "safe-password-123")
        operator_membership = Membership.objects.create(
            organization=self.organization, user=operator, role=Membership.Role.OPERATOR
        )
        CompanyAccessGrant.objects.create(
            organization=self.organization,
            membership=operator_membership,
            company=self.company,
            modules=[ProductModule.Code.NFSE],
            capabilities=["*"],
        )
        self.client.force_login(operator)
        session = self.client.session
        session["hub_organization_id"] = str(self.organization.id)
        session.save()

        response = self.client.get(reverse("hub:nfse-center") + "?view=catalog")

        self.assertContains(response, "Acumuladores por empresa")
        self.assertContains(response, "AC-001")
        self.assertNotContains(response, "AC-999")

        filtered = self.client.get(
            reverse("hub:nfse-center"),
            {"view": "catalog", "catalog_q": "Serviços prestados"},
        )
        empty = self.client.get(
            reverse("hub:nfse-center"),
            {"view": "catalog", "catalog_q": "acumulador inexistente"},
        )

        self.assertContains(filtered, "AC-001")
        self.assertContains(empty, "Nenhum acumulador corresponde aos filtros")

    def test_nfse_demo_catalog_does_not_offer_a_refused_accumulator_form(self) -> None:
        self.organization.is_demo = True
        self.organization.save(update_fields=["is_demo", "updated_at"])

        response = self.client.get(reverse("hub:nfse-center") + "?view=catalog")

        self.assertContains(response, "Consulta na demonstração")
        self.assertNotContains(response, 'value="add_accumulator_rule"')

    def test_nfse_export_history_paginates_without_hiding_old_packages(self) -> None:
        document, _, _ = create_document_and_artifact(
            company=self.company, original_xml="<nfse id='history-page' />", normalized_data={}
        )
        exports = NfseExport.objects.bulk_create(
            [
                NfseExport(
                    organization=self.organization,
                    content_hash=f"{index:064x}",
                    document_count=1,
                    snapshot={"documents": [{"document_id": str(document.pk)}]},
                    target="conference_only_pending_dominio_layout",
                )
                for index in range(101)
            ]
        )
        for export in exports:
            export.documents.add(document)

        first = self.client.get(reverse("hub:nfse-center") + "?view=exports")
        second = self.client.get(reverse("hub:nfse-center") + "?view=exports&exports_page=2")

        self.assertContains(first, "Página 1 de 3")
        self.assertContains(first, "Próxima")
        self.assertContains(second, "Página 2 de 3")
        self.assertContains(second, "Anterior")

    def test_nfse_export_download_recovers_from_a_missing_private_file(self) -> None:
        document, artifact, review = create_document_and_artifact(
            company=self.company,
            original_xml=(
                "<NFSe><infNFSe><numero>missing</numero>"
                "<valores><vLiq>1</vLiq></valores></infNFSe></NFSe>"
            ),
            normalized_data={"service_code": "1401"},
        )
        self.assertIsNotNone(document)
        self.assertIsNone(artifact)
        self.assertIsNotNone(review)
        assert document is not None
        artifact = IntegrationArtifact.objects.create(
            organization=self.organization,
            document=document,
            accumulator_code="AC-MISSING",
            applied_rule="test",
            confidence=100,
        )
        export = create_nfse_export(
            organization=self.organization, documents=[document], actor=self.user
        )
        with patch("django.db.models.fields.files.FieldFile.open", side_effect=OSError):
            failed = self.client.post(reverse("hub:nfse-export-download", args=[export.id]))
        self.assertRedirects(failed, reverse("hub:nfse-center") + "?view=exports")
        export.refresh_from_db()
        self.assertEqual(export.state, NfseExport.State.READY)
        self.assertIsNone(export.downloaded_at)
        export.content.delete(save=False)

        response = self.client.post(reverse("hub:nfse-export-download", args=[export.id]))

        self.assertRedirects(response, reverse("hub:nfse-center") + "?view=exports")
        export.refresh_from_db()
        self.assertEqual(export.state, NfseExport.State.READY)

    def test_nfse_catalog_can_add_an_accumulator_for_future_classifications(self) -> None:
        response = self.client.post(
            reverse("hub:nfse-center"),
            {
                "action": "add_accumulator_rule",
                "company_id": str(self.company.id),
                "accumulator_code": "AC-FUTURO",
                "name": "Novo acumulador",
            },
        )

        self.assertRedirects(response, reverse("hub:nfse-center") + "?view=catalog")
        self.assertTrue(
            AccumulatorRule.objects.filter(
                organization=self.organization,
                company=self.company,
                accumulator_code="AC-FUTURO",
                active=True,
            ).exists()
        )
        self.assertTrue(
            AccumulatorHistoryEntry.objects.filter(
                organization=self.organization,
                company=self.company,
                accumulator_code="AC-FUTURO",
                source=AccumulatorHistoryEntry.Source.MANUAL,
            ).exists()
        )
        catalog = self.client.get(reverse("hub:nfse-center") + "?view=catalog")
        self.assertContains(catalog, "AC-FUTURO")
        self.assertContains(catalog, "Cadastro manual")

        document, artifact, review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='catalog-only' />",
            normalized_data={"service_code": "unknown"},
        )
        self.assertIsNotNone(document)
        self.assertIsNone(artifact)
        self.assertIsNotNone(review)
        assert review is not None
        self.assertEqual(review.suggested_accumulator, "")

    def test_nfse_export_requires_download_before_human_import_confirmation(self) -> None:
        AccumulatorRule.objects.create(
            organization=self.organization,
            company=self.company,
            name="Exportação",
            accumulator_code="AC-EXPORT",
            match={"service_code": "1401"},
        )
        document, artifact, review = create_document_and_artifact(
            company=self.company,
            original_xml=(
                "<NFSe><infNFSe><numero>export</numero>"
                "<valores><vLiq>1</vLiq></valores></infNFSe></NFSe>"
            ),
            normalized_data={"service_code": "1401"},
        )
        self.assertIsNotNone(document)
        self.assertIsNone(review)
        assert document is not None
        if artifact is None:
            artifact = IntegrationArtifact.objects.create(
                organization=self.organization,
                document=document,
                accumulator_code="AC-EXPORT",
                applied_rule="Teste",
                confidence=100,
            )

        created = self.client.post(
            reverse("hub:nfse-center"),
            {"action": "export_dominio_xml", "documents": [str(document.id)]},
        )
        export = NfseExport.objects.get(organization=self.organization)
        exports_page = self.client.get(reverse("hub:nfse-center") + "?view=exports")
        before_download = self.client.post(
            reverse("hub:nfse-export-confirm-import", args=[export.id])
        )
        export.refresh_from_db()
        self.assertEqual(export.state, NfseExport.State.DOWNLOADED)
        downloaded = self.client.post(reverse("hub:nfse-export-download", args=[export.id]))
        export.refresh_from_db()
        confirmed = self.client.post(reverse("hub:nfse-export-confirm-import", args=[export.id]))
        export.refresh_from_db()

        self.assertEqual(created.status_code, 200)
        self.assertEqual(created["Content-Type"], "application/zip")
        self.assertIn("attachment", created["Content-Disposition"])
        self.assertContains(exports_page, "Downloads preparados")
        self.assertContains(exports_page, "Baixar novamente")
        self.assertContains(exports_page, "Código")
        self.assertNotContains(exports_page, "conference_only_pending_dominio_layout")
        self.assertEqual(before_download.status_code, 403)
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded["X-Content-Type-Options"], "nosniff")
        self.assertEqual(confirmed.status_code, 403)
        self.assertEqual(export.state, NfseExport.State.DOWNLOADED)

    def test_review_queue_searches_the_portfolio_and_defaults_to_open_cases(self) -> None:
        _document, _artifact, open_review = create_document_and_artifact(
            company=self.company,
            original_xml="<nfse id='open-filter' />",
            normalized_data={"service_code": "1401"},
        )
        assert open_review is not None
        resolved_company = ClientCompany.objects.create(
            organization=self.organization,
            name="Empresa já conferida",
            dominio_code="0777",
        )
        _document, _artifact, resolved_review = create_document_and_artifact(
            company=resolved_company,
            original_xml="<nfse id='resolved-filter' />",
            normalized_data={"service_code": "1702"},
        )
        assert resolved_review is not None
        resolved_review.status = ReviewCase.Status.RESOLVED
        resolved_review.resolved_accumulator = "AC-777"
        resolved_review.resolved_by = self.user
        resolved_review.resolved_at = timezone.now()
        resolved_review.save()
        IntegrationArtifact.objects.create(
            organization=self.organization,
            document=resolved_review.document,
            accumulator_code="AC-777",
            applied_rule="Decisão humana",
            confidence=100,
        )

        old_queue = self.client.get(reverse("hub:reviews"))
        self.assertRedirects(old_queue, reverse("hub:nfse-center") + "?status=unclassified")
        queue = self.client.get(reverse("hub:nfse-center"), {"status": "unclassified"})
        self.assertContains(queue, self.company.name)
        self.assertNotContains(queue, resolved_company.name)

    def test_dashboard_exposes_exact_operational_queues(self) -> None:
        for code in (
            ProductModule.Code.INTEGRA,
            ProductModule.Code.GUIDES,
            ProductModule.Code.TRIAGE,
            ProductModule.Code.RECONCILIATION,
        ):
            ProductModule.objects.update_or_create(
                organization=self.organization,
                code=code,
                defaults={"enabled": True, "enabled_at": timezone.now()},
            )
        response = self.client.get(reverse("hub:dashboard"), {"view": "portfolio"})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Pendências por área")
        self.assertContains(response, f"{reverse('hub:nfse-center')}?status=unclassified")
        self.assertContains(response, f"{reverse('hub:dte-center')}?status=unread")
        self.assertContains(response, f"{reverse('hub:guides')}?status=pending")

    def test_dashboard_prioritizes_operational_activity_states_without_merging_them(self) -> None:
        today = timezone.localdate()
        overdue = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="ATRASADA",
            title="Atividade em atraso",
            area="fiscal",
            internal_due_on=today - timedelta(days=1),
        )
        due_today = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="HOJE",
            title="Atividade para hoje",
            area="accounting",
            internal_due_on=today,
        )
        unavailable = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="FONTE-INDISPONIVEL",
            title="Fonte indisponível",
            area="payroll",
            freshness=OperationalActivity.Freshness.UNAVAILABLE,
        )

        dashboard = self.client.get(reverse("hub:dashboard"), {"view": "portfolio"})
        source_queue = self.client.get(reverse("hub:activities"), {"freshness": "unavailable"})
        today_queue = self.client.get(reverse("hub:activities"), {"due": "today"})

        self.assertEqual(dashboard.status_code, 200)
        self.assertContains(dashboard, "Em atraso")
        self.assertContains(dashboard, "Para hoje")
        self.assertContains(dashboard, "Fonte indisponível")
        self.assertContains(
            dashboard, f"{reverse('hub:dashboard')}?view=portfolio&amp;filter=overdue"
        )
        self.assertContains(
            dashboard, f"{reverse('hub:dashboard')}?view=portfolio&amp;filter=today"
        )
        self.assertContains(
            dashboard, f"{reverse('hub:dashboard')}?view=portfolio&amp;filter=unavailable"
        )
        self.assertContains(source_queue, unavailable.title)
        self.assertNotContains(source_queue, overdue.title)
        self.assertContains(today_queue, due_today.title)
        self.assertNotContains(today_queue, overdue.title)

    def test_administrator_sees_work_distribution_and_can_open_exact_assignment_queue(self) -> None:
        operator = User.objects.create_user("operator-workload@example.test", "safe-password-123")
        Membership.objects.create(
            organization=self.organization, user=operator, role=Membership.Role.OPERATOR
        )
        assigned = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="ASSIGNED-WORK",
            title="Atividade atribuida",
            area="fiscal",
            assigned_to=operator,
        )
        unassigned = OperationalActivity.objects.create(
            organization=self.organization,
            company=self.company,
            code="UNASSIGNED-WORK",
            title="Atividade sem responsavel",
            area="fiscal",
        )

        dashboard = self.client.get(reverse("hub:dashboard"), {"view": "management"})
        assigned_queue = self.client.get(reverse("hub:activities"), {"assignee": operator.id})
        unassigned_queue = self.client.get(reverse("hub:activities"), {"unassigned": "1"})

        self.assertEqual(dashboard.status_code, 200)
        self.assertContains(dashboard, "Distribuição de atividades")
        self.assertContains(dashboard, operator.email)
        self.assertContains(dashboard, f"{reverse('hub:activities')}?assignee={operator.id}")
        self.assertContains(dashboard, "1 sem responsável")
        self.assertContains(assigned_queue, assigned.title)
        self.assertNotContains(assigned_queue, unassigned.title)
        self.assertContains(unassigned_queue, unassigned.title)
        self.assertNotContains(unassigned_queue, assigned.title)

    def test_review_detail_links_to_exact_case_and_scopes_original_xml(self) -> None:
        original = "<nfse id='case-owned' />"
        _document, _artifact, review = create_document_and_artifact(
            company=self.company,
            original_xml=original,
            normalized_data={
                "number": "NFS-2026-00042",
                "issued_at": "2026-09-21T14:30:00-03:00",
                "service_code": "1401",
                "service_description": "Assessoria contábil mensal",
                "amount": "1234.5",
                "counterparty_ref": "contraparte-pseudonimizada",
            },
        )
        assert review is not None
        detail_url = reverse("hub:review-detail", args=[review.id])
        xml_url = reverse("hub:review-original-xml", args=[review.id])

        dashboard = self.client.get(reverse("hub:dashboard"), {"view": "portfolio"})
        detail = self.client.get(detail_url)
        downloaded = self.client.get(xml_url)

        self.assertContains(dashboard, reverse("hub:nfse-center") + "?status=unclassified")
        center = self.client.get(reverse("hub:nfse-center"), {"status": "unclassified"})
        self.assertContains(center, "1401")
        self.assertContains(center, "NFS-2026-00042")
        self.assertContains(center, "Assessoria contábil mensal")
        self.assertContains(center, detail_url)
        self.assertContains(detail, "Código de serviço")
        self.assertContains(detail, "1401")
        self.assertContains(detail, "Número da NFS-e")
        self.assertContains(detail, "NFS-2026-00042")
        self.assertContains(detail, "Emissão / competência")
        self.assertContains(detail, "21/09/2026")
        self.assertContains(detail, "Descrição do serviço")
        self.assertContains(detail, "Assessoria contábil mensal")
        self.assertContains(detail, "Valor do serviço")
        self.assertContains(detail, "R$ 1.234,50")
        self.assertContains(detail, "Referência da contraparte")
        self.assertContains(detail, "contraparte-pseudonimizada")
        self.assertContains(detail, xml_url)
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.content.decode(), original)
        self.assertIn("attachment", downloaded["Content-Disposition"])
        self.assertEqual(downloaded["Cache-Control"], "private, no-store")

        other_office = Organization.objects.create(
            name="Outro escritório", slug="review-other-office"
        )
        other_company = ClientCompany.objects.create(
            organization=other_office, name="Outra empresa", dominio_code="899"
        )
        _other_document, _other_artifact, other_review = create_document_and_artifact(
            company=other_company,
            original_xml="<nfse id='case-other' />",
            normalized_data={},
        )
        assert other_review is not None
        self.assertEqual(
            self.client.get(reverse("hub:review-detail", args=[other_review.id])).status_code,
            404,
        )
        self.assertEqual(
            self.client.get(reverse("hub:review-original-xml", args=[other_review.id])).status_code,
            404,
        )

    @override_settings(
        CICA_REPORTING_URL="http://127.0.0.1:3080", CICA_REPORTING_SHARED_SECRET="test-secret"
    )
    def test_financial_reports_are_scoped_queued_and_audited(self) -> None:
        source = AccountingBalanceSnapshot.objects.create(
            organization=self.organization,
            company=self.company,
            competence=date(2026, 9, 1),
            source_kind=AccountingBalanceSnapshot.SourceKind.IMPORT,
            source_reference="balancete-2026-09",
        )
        DreMappingSet.objects.create(
            organization=self.organization,
            version=1,
            label="Mapa ativo",
            is_active=True,
            created_by=self.user,
        )
        url = reverse(
            "hub:company-financial-report",
            args=[self.company.id, "dre", source.id, "pdf"],
        )

        self.assertEqual(self.client.get(url).status_code, 405)
        with (
            self.captureOnCommitCallbacks(execute=True),
            patch("apps.hub.tasks.process_financial_report_export.delay") as queued,
        ):
            response = self.client.post(url)

        self.assertRedirects(
            response,
            reverse("hub:company-detail", args=[self.company.id]) + "#relatorios-financeiros",
            fetch_redirect_response=False,
        )
        export = FinancialReportExport.objects.get(organization=self.organization)
        self.assertEqual(export.state, FinancialReportExport.State.WAITING)
        self.assertEqual(export.snapshot_sha256.__len__(), 64)
        queued.assert_called_once_with(str(export.id))
        event = AuditEvent.objects.get(action="hub.financial_report.queued")
        self.assertEqual(event.metadata["resource"], "dre")
        self.assertEqual(event.metadata["source_id"], str(source.id))
        self.assertEqual(
            set(event.metadata), {"resource", "format", "source_id", "snapshot_sha256"}
        )

        foreign_office = Organization.objects.create(name="Outra", slug="financeiro-outra")
        foreign_company = ClientCompany.objects.create(
            organization=foreign_office, name="Estrangeira"
        )
        foreign_source = AccountingBalanceSnapshot.objects.create(
            organization=foreign_office,
            company=foreign_company,
            competence=date(2026, 9, 1),
            source_kind=AccountingBalanceSnapshot.SourceKind.MANUAL,
            source_reference="fora-do-escopo",
        )
        foreign_url = reverse(
            "hub:company-financial-report",
            args=[foreign_company.id, "dre", foreign_source.id, "pdf"],
        )
        self.assertEqual(self.client.post(foreign_url).status_code, 404)

    def test_financial_report_is_unavailable_without_the_node_renderer_or_dre_mapping(self) -> None:
        source = AccountingBalanceSnapshot.objects.create(
            organization=self.organization,
            company=self.company,
            competence=date(2026, 9, 1),
            source_kind=AccountingBalanceSnapshot.SourceKind.MANUAL,
            source_reference="confirmacao",
        )
        url = reverse(
            "hub:company-financial-report",
            args=[self.company.id, "dre", source.id, "xlsx"],
        )
        self.assertEqual(self.client.post(url).status_code, 409)
        DreMappingSet.objects.create(
            organization=self.organization,
            version=1,
            label="Mapa ativo",
            is_active=True,
            created_by=self.user,
        )
        self.assertEqual(self.client.post(url).status_code, 503)

    def test_the_anonymous_workspace_is_blocked(self) -> None:
        self.client.logout()
        anonymous = self.client.get(reverse("hub:dashboard"))
        self.assertEqual(anonymous.status_code, 302)
        self.assertIn(reverse("hub:login"), anonymous["Location"])
